#!/usr/bin/env python3
"""Check that an exported .xlsx is readable by something other than Excel.

    python scripts/verify-cached-values.py <workbook.xlsx>

A formula cell in OOXML carries both the formula (<f>) and its last computed
result (<v>). Excel recalculates on open, so a missing or stale <v> is
invisible there; openpyxl, pandas and SheetJS read the <v> and nothing else.

Check 1  every formula cell carries a cached result.
Check 2  the cached results agree with the inputs they are computed from,
         which catches the nastier variant: a value that is present but left
         over from before the inputs existed.

Check 1 is read off the sheet XML rather than through openpyxl on purpose. A
formula whose answer is the empty string is stored as <v/>, and openpyxl
reports that as None — indistinguishable from a formula that was never
computed. Only the presence of the <v> element separates the two, so the
openpyxl-level view is reported separately, as a count.
"""

import re
import sys
import zipfile

import openpyxl

CELL = re.compile(r"<c\b([^>]*?)(?:/>|>([\s\S]*?)</c>)")
FORMULA = re.compile(r"<f\b[^>]*?(?:/>|>[\s\S]*?</f>)")
VALUE = re.compile(r"<v\b[^>]*?(?:/>|>([\s\S]*?)</v>)")
REF = re.compile(r'r="([^"]+)"')
SHEET = re.compile(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"')
REL = re.compile(r'<Relationship Id="([^"]+)"[^>]*Target="([^"]+)"')


def sheet_parts(zf):
    rels = dict(REL.findall(zf.read("xl/_rels/workbook.xml.rels").decode("utf8")))
    for name, rid in SHEET.findall(zf.read("xl/workbook.xml").decode("utf8")):
        target = rels[rid].lstrip("/")
        yield name, target if target.startswith("xl/") else "xl/" + target


def check_one(path):
    """Check 1: every <f> has a <v>."""
    missing, empty, total = [], 0, 0
    with zipfile.ZipFile(path) as zf:
        for name, part in sheet_parts(zf):
            xml = zf.read(part).decode("utf8")
            for m in CELL.finditer(xml):
                body = m.group(2)
                if not body or not FORMULA.search(body):
                    continue
                total += 1
                ref = REF.search(m.group(1))
                v = VALUE.search(body)
                if v is None:
                    missing.append(f"{name}!{ref.group(1) if ref else '?'}")
                elif not (v.group(1) or ""):
                    empty += 1
    return total, missing, empty


def check_two(path):
    """Check 2: the Internal Summary's D column agrees with B and C."""
    values = openpyxl.load_workbook(path, data_only=True)
    if "Internal Summary" not in values.sheetnames:
        return []
    ws = values["Internal Summary"]
    stale = []
    for row in range(3, 29):
        listed = ws[f"B{row}"].value
        discount = ws[f"C{row}"].value or 0
        customer = ws[f"D{row}"].value
        if not isinstance(listed, (int, float)) or not listed:
            continue
        if not isinstance(customer, (int, float)) or not isinstance(discount, (int, float)):
            continue
        expected = listed - (listed * discount)
        if abs(customer - expected) > 0.01:
            stale.append(f"Internal Summary!D{row}: cached {customer:,.2f}, expected {expected:,.2f}")
    return stale


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = sys.argv[1]
    ok = True

    total, missing, empty = check_one(path)
    print(f"check 1: {len(missing)} of {total} formula cells have no cached value")
    for m in missing[:20]:
        print("   ", m)
    if len(missing) > 20:
        print(f"    … and {len(missing) - 20} more")
    print(f"         ({empty} more are computed and genuinely empty - an IF(...,\"\") answer)")
    ok &= not missing

    stale = check_two(path)
    print(f"check 2: {len(stale)} cached values disagree with their inputs")
    for s in stale:
        print("   ", s)
    ok &= not stale

    print("OK" if ok else "FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
