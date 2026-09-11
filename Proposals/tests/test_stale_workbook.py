"""A workbook whose cached results predate its inputs must be refused.

An operator exported an RFC from a separate program, dropped it in, and got a
proposal with the financing sections missing. Financing was not the problem:

    INPUT SHEET!E8 = 'TP5-120-480-1' x 6   <- typed in, plainly visible
    INPUT SHEET!K8 (line total)            <- formula, never computed
      -> Internal Summary!D30 = 0
      -> Financial Worksheet!B44 = 0
      -> DLL Schedule!D5 (Loan_Amount) = 0
      -> has_financing = False, Sections 17 / 17A-D suppressed

Every cost printed $0.00 and the document rendered anyway. The file is FULL of
numbers, so nothing looked wrong - which is what makes this worse than an
obviously empty workbook.

The detection is semantic, and that matters. The obvious signal -
`fullCalcOnLoad="1"` in the package - is useless: **openpyxl sets it on every
save**, so it fires on any programmatically written workbook including this
project's own synthetic fixture, which is correctly valued and must load.
These tests pin that distinction, because getting it wrong blocks real work.
"""

from __future__ import annotations

import pytest

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.engine import stale_cached_values
from ev_proposal_agent.errors import WorkbookNotCalculated

from .conftest import (BEST_WESTERN, LEGACY, MARRIOTT, MARRIOTT_L2, ROOT,
                       SHOWCASE, V16)

EXPORTS = [
    pytest.param(ROOT / "inputs" / "project-rfc-msrp-calculator (1).xlsx", id="export-1"),
    pytest.param(ROOT / "inputs" / "project-rfc-msrp-calculator (2).xlsx", id="export-2"),
]

REAL = [
    pytest.param(BEST_WESTERN, id="v16"),
    pytest.param(V16, id="v16-master"),
    pytest.param(SHOWCASE, id="v16-legacy-engine"),
    pytest.param(MARRIOTT, id="v15_no_itc"),
    pytest.param(MARRIOTT_L2, id="v16-l2-only"),
    pytest.param(LEGACY, id="legacy_food4less"),
]


@pytest.mark.parametrize("path", EXPORTS)
def test_a_stale_export_is_refused(path, field_map):
    if not path.exists():
        pytest.skip(f"missing specimen: {path}")
    with pytest.raises(WorkbookNotCalculated) as excinfo:
        with engine_mod.load(path, field_map):
            pass
    err = excinfo.value
    assert "Ctrl+Alt+F9" in str(err), "the operator must be told what to do"
    # The detail names the actual cell, so the report is diagnosable rather
    # than just a complaint.
    assert "INPUT SHEET!E8" in (err.detail or ""), err.detail


@pytest.mark.parametrize("path", REAL)
def test_real_workbooks_load_untouched(path, field_map):
    with engine_mod.load(path, field_map) as lw:
        assert lw.detection.chassis
        assert not stale_cached_values(lw.values)


def test_the_check_is_semantic_not_a_package_flag(tmp_path):
    """An openpyxl round-trip must not turn a good workbook into a bad one.

    openpyxl stamps `fullCalcOnLoad="1"` on save. If the guard keyed off that,
    re-saving any valid workbook would make the app refuse it - and the
    synthetic three-tier fixture in `tests/fixtures.py` is built exactly that
    way.
    """
    import zipfile

    from openpyxl import load_workbook

    resaved = tmp_path / "resaved.xlsx"
    load_workbook(BEST_WESTERN).save(resaved)

    book = zipfile.ZipFile(resaved).read("xl/workbook.xml").decode("utf-8", "ignore")
    assert 'fullCalcOnLoad="1"' in book, (
        "openpyxl no longer sets this - the comment explaining why the guard "
        "is semantic needs revisiting, but the guard itself is still correct")

    # openpyxl drops cached values on save, so this copy IS genuinely
    # unusable - but it must be unusable for a stated reason, not because of
    # a package flag.
    values = load_workbook(resaved, data_only=True)
    problems = stale_cached_values(values)
    assert all("fullCalcOnLoad" not in p for p in problems)


def test_a_stale_workbook_names_the_cells_that_disagree(field_map):
    path = ROOT / "inputs" / "project-rfc-msrp-calculator (2).xlsx"
    if not path.exists():
        pytest.skip(f"missing specimen: {path}")
    from openpyxl import load_workbook

    problems = stale_cached_values(load_workbook(path, data_only=True))
    assert problems, "the export is stale and the checker found nothing"
    joined = " ".join(problems)
    assert "line total" in joined
    # Internal Summary rows whose inputs are filled but whose customer price
    # is still zero - D9 is the clearest: B9=4400, C9=7%, D9 cached 0.
    assert any("Internal Summary!D" in p for p in problems), problems


def test_a_healthy_workbook_reports_no_discrepancies():
    from openpyxl import load_workbook

    assert stale_cached_values(load_workbook(BEST_WESTERN, data_only=True)) == []
