"""Phases 5 to 8 end to end: numbering, charts, narrative, validation, render."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from docx import Document

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.extract import build_context
from ev_proposal_agent.narrative import build_all, verify_numerals
from ev_proposal_agent.render import generate, safe_filename
from ev_proposal_agent.sections import build_plan, numbering_is_consistent
from ev_proposal_agent.validate import validate

from .conftest import SHOWCASE_OPERATOR_INPUTS

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "proposal_template.docx"

pytestmark = pytest.mark.skipif(
    not TEMPLATE.exists(), reason="run `python -m tools.build_template` first"
)


# --------------------------------------------------------------------------
# section numbering
# --------------------------------------------------------------------------


def _plan(*, disabled=(), **flags):
    base = {"has_history": True, "has_baseline": True, "has_financing": True}
    base.update(flags)
    return build_plan(base, disabled=disabled,
                      substitutions={"projection_years": 5,
                                     "loan_term_months": 60})


def test_full_document_numbering():
    plan = _plan()
    assert numbering_is_consistent(plan) == []
    numbers = [r.number for r in plan.sections]
    assert numbers[:5] == ["1", "2", "3", "3A", "3B"]
    assert numbers[-1] == "18"


def test_no_history_renumbers_everything_after_it():
    """Sections 3, 3A and 3B go; what was 4 becomes 3 and the tail shifts."""
    plan = _plan(has_history=False, has_baseline=False)
    assert numbering_is_consistent(plan) == []
    assert not plan.is_shown("history")
    assert not plan.is_shown("history_actuals")
    assert plan.number_of("methodology") == "3"
    assert plan.number_of("assumptions") == "3A"
    # 18 top-level sections normally; losing Section 3 makes the last one 17.
    assert plan.number_of("terms") == "17"


def test_no_financing_drops_five_sections():
    plan = _plan(has_financing=False)
    assert numbering_is_consistent(plan) == []
    for key in ("financing", "financing_monthly", "financing_table_1",
                "financing_table_2", "financing_table_3"):
        assert not plan.is_shown(key)
    assert plan.number_of("scope_of_work") == "16"
    assert plan.number_of("terms") == "17"


def test_baseline_can_drop_without_taking_section_3_with_it():
    """History present but no port counts: 3 and 3A stay, 3B goes."""
    plan = _plan(has_baseline=False)
    assert plan.is_shown("history_actuals")
    assert not plan.is_shown("history_baseline")
    assert numbering_is_consistent(plan) == []


def test_continuation_goes_when_its_own_requirement_is_unmet():
    """3B declares both has_history and has_baseline, so losing history drops it
    on its own requirement before the parent check is ever reached."""
    plan = _plan(has_history=False, has_baseline=True)
    assert not plan.is_shown("history_baseline")
    reasons = {s.key: why for s, why in plan.dropped}
    assert "has_history" in reasons["history_baseline"]


def test_continuation_cannot_outlive_its_parent():
    """The defensive case: a continuation whose own requirements are satisfied
    but whose parent was dropped. No registry entry reaches this today, so it is
    exercised directly to keep the guard honest."""
    from ev_proposal_agent.sections import Section, build_plan

    registry = (
        Section("parent_gone", "PARENT", requires=("never_true",)),
        Section("orphan", "ORPHAN", parent="parent_gone", suffix="A"),
    )
    import ev_proposal_agent.sections as sections_mod

    original = sections_mod.REGISTRY
    try:
        sections_mod.REGISTRY = registry
        plan = build_plan({"never_true": False})
    finally:
        sections_mod.REGISTRY = original

    assert not plan.is_shown("orphan")
    assert any("parent" in why for _s, why in plan.dropped)


def test_neither_history_nor_financing():
    plan = _plan(has_history=False, has_baseline=False, has_financing=False)
    assert numbering_is_consistent(plan) == []
    assert len(plan.sections) == 28 - 3 - 5


@pytest.mark.parametrize("flags", [
    {}, {"has_history": False, "has_baseline": False},
    {"has_financing": False},
    {"has_history": False, "has_baseline": False, "has_financing": False},
])
def test_all_three_listings_agree(flags):
    """The badge, the page-2 contents and the Section 2 grid come from one plan,
    so they cannot disagree - this pins that down."""
    plan = _plan(**flags)
    toc = {row["number"] for row in plan.toc_rows()}
    badges = {r.number for r in plan.sections}
    assert toc == badges

    listed = {pair["number"] for pair in plan.contents_pairs()}
    listed |= {pair["number2"] for pair in plan.contents_pairs() if pair["number2"]}
    top_level = {r.number for r in plan.sections if r.in_contents}
    assert listed == top_level


def test_no_number_is_skipped():
    plan = _plan(has_history=False, has_baseline=False)
    tops = [int(r.number) for r in plan.sections if r.number.isdigit()]
    assert tops == list(range(1, len(tops) + 1))


# --------------------------------------------------------------------------
# narrative
# --------------------------------------------------------------------------


def test_narrative_uses_only_extracted_numbers(showcase_context):
    """Every numeral in every generated sentence must trace to the workbook."""
    for token, text in build_all(showcase_context).items():
        if not text:
            continue
        unknown = verify_numerals(text, showcase_context)
        assert not unknown, f"{token} invented {unknown}"


def test_narrative_adapts_to_a_site_with_no_dc_fast(showcase_context):
    text = build_all(showcase_context)["narrative_site_history"]
    assert "no DC fast charging history" in text


def test_breakeven_wording_flips_when_there_is_none(showcase_context):
    ctx = showcase_context
    original = ctx.raw.get("breakeven_year")
    try:
        ctx.raw["breakeven_year"] = None
        text = build_all(ctx)["narrative_investment"]
        assert "does not turn positive" in text
    finally:
        ctx.raw["breakeven_year"] = original


def test_itc_sentence_absent_when_itc_is_zero(showcase_context):
    text = build_all(showcase_context)["narrative_roi"]
    assert "tax credit" not in text.lower()


def test_verify_numerals_catches_an_invented_figure(showcase_context):
    assert verify_numerals("Utilisation rose by 47.3% last quarter.",
                           showcase_context) == ["47.3"]


# --------------------------------------------------------------------------
# validation
# --------------------------------------------------------------------------


def test_clean_workbook_produces_no_errors(showcase_context, field_map):
    plan = _plan()
    findings = validate(showcase_context, field_map, plan)
    assert findings.errors == []


def test_broken_subtotal_is_an_error(field_map, showcase_path):
    """A cost table that does not foot must never reach a customer.

    Built on its own context rather than the shared fixture - mutating a
    session-scoped context leaks the injected fault into every later test.
    """
    with engine_mod.load(showcase_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=SHOWCASE_OPERATOR_INPUTS)
    ctx.raw["cost_charger_hardware"] = ctx.raw["cost_charger_hardware"] + 5000
    findings = validate(ctx, field_map, _plan())
    assert any(f.code == "cost-subtotal-mismatch" for f in findings.errors)


def test_validate_does_not_mutate_the_context(showcase_context, field_map):
    """Two calls must not double up. The GUI re-validates after every edit."""
    before = len(showcase_context.findings)
    first = validate(showcase_context, field_map, _plan())
    second = validate(showcase_context, field_map, _plan())
    assert len(showcase_context.findings) == before
    assert len(first) == len(second)


# --------------------------------------------------------------------------
# charts
# --------------------------------------------------------------------------


def test_charts_render(field_map, showcase_path, tmp_path):
    from ev_proposal_agent.charts import render_all

    with engine_mod.load(showcase_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=SHOWCASE_OPERATOR_INPUTS)
        made = render_all(lw, ctx, tmp_path)
    assert set(made) == {"chart_historical_revenue", "chart_annual_operating_profit",
                         "chart_cumulative_cashflow"}
    for path in made.values():
        assert path.exists() and path.stat().st_size > 5_000


def test_offset_defined_name_is_resolved(field_map, v16_paths):
    """`cumCashCats` is an OFFSET formula openpyxl cannot evaluate."""
    from ev_proposal_agent.charts import resolve_offset_name

    with engine_mod.load(v16_paths, field_map) as lw:
        cats = resolve_offset_name(lw, "cumCashCats")
    assert cats == [0, 1, 2, 3, 4, 5]      # year 0 plus the 5-year horizon


def test_legacy_chassis_falls_back_when_names_are_absent(field_map, legacy_path):
    from ev_proposal_agent.charts import resolve_offset_name

    with engine_mod.load(legacy_path, field_map) as lw:
        assert resolve_offset_name(lw, "cumCashCats") is None


def test_v15_chassis_falls_back_when_names_are_absent(field_map, marriott_path):
    """Keyed on the name being absent, not on the chassis, so a third layout
    needed no new branch - but it is worth proving that."""
    from ev_proposal_agent.charts import resolve_offset_name

    with engine_mod.load(marriott_path, field_map) as lw:
        assert resolve_offset_name(lw, "cumCashCats") is None


# --------------------------------------------------------------------------
# full render
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def rendered(showcase_path, tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("render") / "Showcase.docx"
    generate(showcase_path, output=out,
             operator_inputs={**SHOWCASE_OPERATOR_INPUTS,
                              "client_contact_name": "Marisol Enriquez",
                              "prepared_by_name": "Paul Arms",
                              "site_name": "Showcase Liquor Pasadena",
                              "site_location_narrative": "On North Lake Avenue."},
             progress=lambda _m: None)
    return out


def test_render_produces_a_document_and_a_qa_report(rendered):
    assert rendered.exists()
    qa = rendered.with_name(rendered.stem + "_QA.txt")
    assert qa.exists()
    text = qa.read_text(encoding="utf-8")
    assert "DETECTED" in text and "VALUES" in text and "SECTIONS" in text


def test_rendered_document_has_no_leftovers(rendered):
    from ev_proposal_agent.render import verify_output

    assert verify_output(rendered) == []


def test_rendered_document_keeps_its_structure(rendered):
    doc = Document(rendered)
    assert len(doc.sections) == 4
    with zipfile.ZipFile(rendered) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    # 7 originals kept (logos and product photos) plus 3 generated charts
    assert len(media) >= 10


def test_footers_carry_the_new_site(rendered):
    with zipfile.ZipFile(rendered) as z:
        for name in [n for n in z.namelist() if re.match(r"word/footer\d+\.xml", n)]:
            text = z.read(name).decode("utf-8")
            assert "Showcase Liquor Pasadena" in text
            assert "Food4Less" not in text


def test_toc_matches_the_badges(rendered):
    doc = Document(rendered)
    toc = [r.cells[0].text.strip() for r in doc.tables[1].rows[1:]
           if r.cells[0].text.strip()]
    assert toc[:5] == ["1", "2", "3", "3A", "3B"]
    assert "18" in toc


def test_cost_table_is_not_duplicated(rendered):
    """The loop must not re-emit the Grand Total that the fixed tail prints."""
    doc = Document(rendered)
    labels = [r.cells[0].text.strip() for r in doc.tables[28].rows]
    assert labels.count("Grand Total") == 1
    assert labels.count("Total After Discount") == 1


def test_ampersand_survives(rendered):
    """docxtpl defaults to autoescape off, which silently eats a bare `&`."""
    doc = Document(rendered)
    labels = [r.cells[0].text for r in doc.tables[28].rows]
    assert any("EVOLV & Commissioning" in label for label in labels)


def test_negative_currency_is_signed_before_the_symbol(rendered):
    doc = Document(rendered)
    year_zero = doc.tables[31].rows[1]
    assert year_zero.cells[1].text.startswith("-$")


def test_itc_row_absent_when_itc_is_zero(rendered):
    doc = Document(rendered)
    labels = [r.cells[0].text.strip() for r in doc.tables[29].rows]
    assert "Federal ITC (30%)" not in labels
    assert "Net Revenues" in labels


# --------------------------------------------------------------------------
# full render - the v15_no_itc chassis
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def rendered_marriott(marriott_path, tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("render_v15") / "Marriott.docx"
    # force=True: this workbook's Cashflow!G3 divides ten years of carbon by
    # five, so it carries a real ERROR until the sheet is fixed. These tests are
    # about rendering mechanics, not the workbook's arithmetic - see
    # test_v15_no_itc_chassis.py for the assertion that the defect IS reported.
    generate(marriott_path, output=out, force=True,
             operator_inputs={"client_contact_name": "Dana Whitfield",
                              "prepared_by_name": "Paul Arms",
                              "site_name": "Marriott Bakersfield",
                              "site_location_narrative": "Off Highway 99."},
             progress=lambda _m: None)
    return out


def test_v15_renders_end_to_end(rendered_marriott):
    from ev_proposal_agent.render import verify_output

    assert rendered_marriott.exists()
    assert rendered_marriott.with_name(
        rendered_marriott.stem + "_QA.txt").exists()
    assert verify_output(rendered_marriott) == []


def test_v15_prints_no_itc_row(rendered_marriott):
    """There is no ITC row on this chassis at all. Reading it off B5 - which
    holds EVSE Revenues - would have printed $416,761 as a federal credit."""
    doc = Document(rendered_marriott)
    cells = {r.cells[0].text.strip(): r.cells[-1].text.strip()
             for r in doc.tables[29].rows}
    assert "Federal ITC (30%)" not in cells
    assert "ITC" not in cells
    # $416,761 belongs to EVSE Revenues and to nothing else. Under the legacy
    # chassis it appeared twice: once here, correctly, and once as a federal
    # credit the site is not receiving.
    assert cells["EVSE Revenues"] == "$416,761.14"
    assert [k for k, v in cells.items() if "416,761" in v] == ["EVSE Revenues"]


def test_v15_cost_table_prints_both_totals(rendered_marriott):
    """A discounted workbook: the list total and the customer total differ, and
    Section 7 shows both so the discount is visible."""
    doc = Document(rendered_marriott)
    cells = {r.cells[0].text.strip(): r.cells[-1].text.strip()
             for r in doc.tables[28].rows}
    assert cells["Grand Total"] == "$345,247.91"
    assert cells["Total After Discount"] == "$253,716.82"


# --------------------------------------------------------------------------
# filenames
# --------------------------------------------------------------------------


@pytest.mark.parametrize("raw, expected", [
    ("Showcase Liquor Pasadena", "Showcase Liquor Pasadena"),
    ("A/B: Test*", "AB Test"),
    ("", "Proposal"),
    (None, "Proposal"),
])
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def test_output_name_omits_a_missing_job_number(showcase_context):
    from ev_proposal_agent.render import default_output_path

    name = default_output_path(showcase_context, directory=Path(".")).name
    assert "__" not in name, "a blank job number must not leave a double separator"


# --------------------------------------------------------------------------
# packaging
# --------------------------------------------------------------------------


def test_bundled_paths_resolve_under_pyinstaller(tmp_path, monkeypatch):
    """`--onefile` unpacks to `sys._MEIPASS`. Everything that reads `config/`
    or `templates/` must go through `paths.py`, or the packaged .exe looks
    beside itself and finds nothing."""
    import importlib

    (tmp_path / "config").mkdir()
    (tmp_path / "templates").mkdir()
    monkeypatch.setattr("sys._MEIPASS", str(tmp_path), raising=False)

    import ev_proposal_agent.paths as paths_mod

    reloaded = importlib.reload(paths_mod)
    try:
        assert reloaded.ROOT == tmp_path
        assert reloaded.FIELD_MAP == tmp_path / "config" / "field_map.yaml"
        assert reloaded.PROPOSAL_TEMPLATE == tmp_path / "templates" / "proposal_template.docx"
    finally:
        monkeypatch.delattr("sys._MEIPASS", raising=False)
        importlib.reload(paths_mod)


def test_output_and_settings_paths_are_per_user():
    from ev_proposal_agent.paths import output_dir, settings_path

    assert output_dir().name == "EV Proposals"
    assert output_dir().parent.name == "Documents"
    assert settings_path().name == "recent.json"


# --------------------------------------------------------------------------
# operator customisation
# --------------------------------------------------------------------------


def test_operator_can_switch_sections_off():
    plan = _plan(disabled={"om", "carbon"})
    assert not plan.is_shown("om")
    assert not plan.is_shown("carbon")
    assert numbering_is_consistent(plan) == []


def test_required_sections_cannot_be_switched_off():
    """A proposal without terms or an executive summary is not a proposal, and
    a tick-box should not be able to produce one."""
    plan = _plan(disabled={"terms", "executive_summary", "overview"})
    for key in ("terms", "executive_summary", "overview"):
        assert plan.is_shown(key)
    assert {s.key for s in plan.forced_on} == {"terms", "executive_summary",
                                               "overview"}


def test_switching_sections_off_renumbers_the_rest():
    plan = _plan(disabled={"rip_and_replace"})
    assert plan.number_of("infrastructure") == "8"      # was 9
    assert plan.number_of("terms") == "17"              # was 18
    tops = [int(r.number) for r in plan.sections if r.number.isdigit()]
    assert tops == list(range(1, len(tops) + 1))


def test_every_section_has_a_template_guard():
    """An unguarded section leaves its badge behind pointing at a numbering
    entry that no longer exists, which is a hard Jinja error at render time."""
    import re as _re
    import zipfile as _zip

    from ev_proposal_agent.sections import REGISTRY

    with _zip.ZipFile(TEMPLATE) as z:
        flat = _re.sub(r"<[^>]+>", "", z.read("word/document.xml").decode("utf-8"))
    for section in REGISTRY:
        assert f"if show_{section.key} " in flat, f"{section.key} is unguarded"


def test_removing_a_section_removes_its_body(showcase_path, tmp_path):
    out = tmp_path / "Trimmed.docx"
    generate(showcase_path, output=out,
             operator_inputs={**SHOWCASE_OPERATOR_INPUTS,
                              "disabled_sections": ["om", "carbon"]},
             progress=lambda _m: None)
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "What O&M means" not in text
    assert "two main types of Carbon Credits" not in text
    assert "Non-binding proposal disclaimer" in text        # locked, still there
    # page orientation survives: the landscape appendix break is untouched
    assert len(doc.sections) == 4


def test_custom_cover_photo_is_used(showcase_path, tmp_path):
    from PIL import Image  # noqa: F401  (pillow ships with matplotlib)

    photo = tmp_path / "site.png"
    Image.new("RGB", (1200, 800), (40, 90, 110)).save(photo)

    out = tmp_path / "WithCover.docx"
    generate(showcase_path, output=out,
             operator_inputs={**SHOWCASE_OPERATOR_INPUTS,
                              "cover_photo_path": str(photo)},
             progress=lambda _m: None)
    assert out.exists()


def test_cover_falls_back_to_the_stock_photo():
    from ev_proposal_agent.paths import TEMPLATES_DIR

    assert (TEMPLATES_DIR / "assets" / "cover_default.jpeg").exists()


# --------------------------------------------------------------------------
# Section 13 carbon KPI strip
# --------------------------------------------------------------------------


def test_carbon_kpis_are_linked_to_the_workbook(showcase_context):
    """All three tiles were still the reference site's hardcoded numbers."""
    ctx = showcase_context
    assert ctx.sources["cc_total_kpi"] == "Financial Worksheet!B4"
    assert ctx.num("cc_total_kpi") == pytest.approx(21525.63, abs=0.01)
    assert ctx.formatted["cc_l2_rate"] == "$0.0045"


def test_only_installed_tiers_appear_in_the_multiplier_tile(showcase_context):
    """`FW!D43:J45` lists all six L3 classes. Only the columns with a quantity
    above zero are being bought; printing the rest quotes credits for hardware
    that is not in the proposal. Showcase installs one 60 kW unit."""
    assert showcase_context.formatted["cc_multiplier_lines"] == "$4,300"
    assert showcase_context.raw["cc_active_tier_count"] == 1


def test_multiple_tiers_stack_as_real_line_breaks(field_map, showcase_path,
                                                  tmp_path):
    """Three active tiers must render as three lines, not one run of text.

    A literal newline inside `<w:t>` is a space in Word, and an unstyled
    RichText emits text with no run wrapper at all, which Word discards - the
    tile came out blank. Both were real bugs and only appear above one tier.
    """
    import openpyxl
    import shutil
    import zipfile

    fixture = tmp_path / "multi_tier.xlsx"
    shutil.copyfile(showcase_path, fixture)
    # data_only=True so the saved copy keeps Excel's cached results. openpyxl
    # cannot preserve formulas AND cached values through a save: keep the
    # formulas and every value reads None, which fails 30 validations.
    wb = openpyxl.load_workbook(fixture, data_only=True)
    fw = wb["Financial Worksheet"]
    fw["F44"], fw["I44"] = 2, 1                   # add 120 kW and 240 kW
    wb.save(fixture)
    wb.close()

    out = tmp_path / "MultiTier.docx"
    generate(fixture, output=out, progress=lambda _m: None)

    doc = Document(out)
    tile = next(c for t in doc.tables for r in t.rows for c in r.cells
                if "LEVEL 3 MULTIPLIERS" in c.text)
    assert tile.paragraphs[0].text == "$4,300\n$8,600\n$17,200"

    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    index = xml.find("LEVEL 3 MULTIPLIERS")
    cell = xml[xml[:index].rfind("<w:tc>"):index]
    assert cell.count("<w:br/>") == 2, "line breaks must be <w:br/>, not '\n'"
    # every figure run keeps the tile's teal 16pt bold
    assert cell.count('w:val="08B3AD"') >= 3
