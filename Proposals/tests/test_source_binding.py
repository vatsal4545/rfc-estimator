"""Every printed figure is bound to the cell it claims to come from.

The suite already tests each hop of the chain separately:

    workbook cell  ->  ctx token  ->  template token  ->  rendered position

and never tests the chain. Exactly one existing test compares a workbook cell
to a value on a rendered page (`test_infrastructure_section.py`), and it checks
a sum rather than a token.

Worse, the dominant pattern elsewhere is structurally blind:

    assert ctx.num("cost_grand_total") == pytest.approx(43235.97)

That literal was obtained by running the extractor. If the extractor reads a
plausible-but-wrong cell, the test records the wrong value as correct. That is
how a fictional $416,761 "Federal ITC (30%)" line survived until a human
noticed it.

These tests read the workbook independently of the extractor and compare. A
wrong cell fails here even when it is internally consistent everywhere else.
"""

from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from .conftest import (BEST_WESTERN, LEGACY, MARRIOTT, MARRIOTT_L2, SHOWCASE, V16)

# "Financial Worksheet!B44" - one cell, nothing else. Ranges, arithmetic and
# prose derivations are deliberately excluded: this test verifies the claims it
# can verify exactly, and says how many it skipped.
SINGLE_CELL = re.compile(r"^([^!]+)!([A-Z]{1,3}\d+)$")

SPECIMENS = [
    pytest.param("best_western_context", BEST_WESTERN, id="v16-scenario-grid"),
    pytest.param("v16_context", V16, id="v16-master"),
    pytest.param("showcase_context", SHOWCASE, id="v16-legacy-engine"),
    pytest.param("marriott_context", MARRIOTT, id="v15_no_itc"),
    pytest.param("marriott_l2_context", MARRIOTT_L2, id="v16-l2-only"),
    pytest.param("legacy_context", LEGACY, id="legacy_food4less"),
]

# Tokens whose stored value is deliberately NOT the raw cell. Each is a real
# transform with a reason, verified by its own test elsewhere.
TRANSFORMED = {
    # published as a fraction of the horizon, not the horizon total
    "carbon_per_year",
    # sign is flipped for display: the sheet holds cost as a negative
    "roi_total_costs_upfront",
}


def _cells(path):
    return load_workbook(path, data_only=True, read_only=False)


def _comparable(stored, cell):
    """True when the stored value is the cell's value.

    `None` and `0` are the same fact here: `_coerce_zero` turns a blank cost
    row into 0.0 on the way in, deliberately, because Excel leaves unused rows
    empty rather than zeroed. Whitespace is normalised because the client-info
    block is sanitised on the way in - `INPUT SHEET!N21` carries hard newlines
    that must not reach a Word paragraph as literal breaks.
    """
    if isinstance(stored, bool) or isinstance(cell, bool):
        return bool(stored) == bool(cell)
    if isinstance(stored, (int, float)) and isinstance(cell, (int, float)):
        return abs(float(stored) - float(cell)) < 0.005
    if stored in (None, "", 0, 0.0) and cell in (None, "", 0, 0.0):
        return True
    if isinstance(stored, (int, float)) and cell is None:
        return abs(float(stored)) < 0.005
    return " ".join(str(stored).split()) == " ".join(str(cell).split())


@pytest.mark.parametrize("fixture_name,path", SPECIMENS)
def test_every_emitted_token_records_a_source(fixture_name, path, request):
    """A printed number with no provenance cannot be checked by anyone."""
    ctx = request.getfixturevalue(fixture_name)
    operator = set(ctx.raw.get("_operator_tokens") or ()) | {
        "site_name", "client_contact_name", "prepared_by_name",
        "site_location_narrative", "client_title", "cover_photo",
        "photo_level2", "photo_level3",
    }
    missing = sorted(t for t in ctx.emitted
                     if t not in ctx.sources and t not in operator)
    assert not missing, (
        "these tokens print a value with no recorded source, so nothing can "
        f"trace them back to the workbook: {missing}"
    )


@pytest.mark.parametrize("fixture_name,path", SPECIMENS)
def test_single_cell_tokens_equal_their_own_cell(fixture_name, path, request):
    """Read the cell the token names, independently, and compare.

    This is the hop nothing else covers. It does not care whether the value is
    plausible or internally consistent - only whether it is the cell the
    context says it is.
    """
    ctx = request.getfixturevalue(fixture_name)
    wb = _cells(path)
    checked, skipped, wrong = 0, 0, []
    try:
        for token in sorted(ctx.emitted):
            src = ctx.sources.get(token)
            if not src or token in TRANSFORMED:
                skipped += 1
                continue
            m = SINGLE_CELL.match(src)
            if not m:
                skipped += 1
                continue
            sheet, ref = m.group(1), m.group(2)
            # A source is a CLAIM THAT THE VALUE IS THIS CELL only when it is
            # exactly "Sheet!Ref". Anything with a prefix - "recomputed from
            # DLL Schedule!D9", "detected from the formula of ...!I6" - says
            # the opposite: the cell is where the derivation started, not what
            # the token holds. Those are honest and unverifiable here.
            if sheet not in wb.sheetnames:
                skipped += 1
                continue
            checked += 1
            cell = wb[sheet][ref].value
            stored = ctx.raw.get(token)
            if not _comparable(stored, cell):
                wrong.append(f"{token}: context has {stored!r}, {sheet}!{ref} holds {cell!r}")
    finally:
        wb.close()

    assert not wrong, (
        f"{len(wrong)} token(s) do not match the cell they name:\n  "
        + "\n  ".join(wrong)
    )
    # Guard against the assertion passing vacuously if sources ever stop being
    # cell-shaped.
    assert checked >= 20, (
        f"only {checked} single-cell bindings were checkable on this specimen "
        f"({skipped} skipped) - the test has stopped covering anything"
    )


# ---------------------------------------------------------------------------
# Mutation: the only construct that catches a wrong cell in general.
#
# Every other test compares a token to a value obtained by running the
# extractor, so a token wired to the wrong cell is recorded as correct. Here we
# CHANGE a cell and demand that exactly the tokens claiming it move. A token
# that names B44 but reads B43 does not move when B44 changes, and fails.
# ---------------------------------------------------------------------------

# Perturbing a cached value does not re-run Excel, so dependants keep their old
# cached numbers. That is what makes this a clean isolation test: only tokens
# reading the cell DIRECTLY are expected to move.
DELTA = 1234.5

# Resolved inside `engine.load()` - detection runs once, before any context is
# built - so mutating the cell afterwards cannot reach them. Their binding is
# covered by `test_single_cell_tokens_equal_their_own_cell` instead. This is a
# limit of the harness, not a licence: anything added here must be genuinely
# load-time, and `detect_horizon` reading B10 is exactly that.
DETECTED_AT_LOAD = {"projection_years", "revenue_scenario", "engine",
                    "has_history", "has_financing"}


def _numeric_bindings(ctx, wb):
    """token -> (sheet, ref) for every token that claims one numeric cell."""
    out = {}
    for token in ctx.emitted:
        src = ctx.sources.get(token)
        if not src:
            continue
        m = SINGLE_CELL.match(src)
        if not m or m.group(1) not in wb.sheetnames:
            continue
        if isinstance(ctx.raw.get(token), bool) or token in DETECTED_AT_LOAD:
            continue
        if isinstance(ctx.raw.get(token), (int, float)):
            out[token] = (m.group(1), m.group(2))
    return out


@pytest.mark.parametrize("path", [BEST_WESTERN, LEGACY], ids=["v16", "legacy"])
def test_moving_a_cell_moves_exactly_the_tokens_that_claim_it(path, field_map):
    from ev_proposal_agent import engine as engine_mod
    from ev_proposal_agent.extract import build_context

    inputs = {"site_name": "Site", "client_contact_name": "Contact",
              "prepared_by_name": "Preparer", "site_location_narrative": "N.",
              "existing_ports_l2": 5, "existing_ports_l3": 2}

    wb = _cells(path)
    try:
        with engine_mod.load(path, field_map) as lw:
            base = build_context(lw, field_map, operator_inputs=dict(inputs))
            bindings = _numeric_bindings(base, wb)
            assert len(bindings) >= 30, (
                f"only {len(bindings)} numeric cell bindings found - nothing to mutate"
            )

            # One token per distinct cell, so "everything else must hold still"
            # is not defeated by two tokens legitimately sharing a cell.
            by_cell: dict[tuple[str, str], list[str]] = {}
            for tok, cell in bindings.items():
                by_cell.setdefault(cell, []).append(tok)

            targets = sorted(by_cell)[:12]
            failures: list[str] = []

            for sheet, ref in targets:
                claimants = sorted(by_cell[(sheet, ref)])
                original = lw.values[sheet][ref].value
                if not isinstance(original, (int, float)) or isinstance(original, bool):
                    continue
                lw.values[sheet][ref] = float(original) + DELTA
                try:
                    after = build_context(lw, field_map, operator_inputs=dict(inputs))
                finally:
                    lw.values[sheet][ref] = original

                for tok in claimants:
                    before_v, after_v = base.raw.get(tok), after.raw.get(tok)
                    if not isinstance(after_v, (int, float)):
                        continue
                    if abs(float(after_v) - float(before_v)) < 0.005:
                        failures.append(
                            f"{tok} claims {sheet}!{ref}, but changing that cell by "
                            f"{DELTA} left it at {before_v!r} - it is reading "
                            f"somewhere else"
                        )

                # Nothing bound to a DIFFERENT cell may move.
                for tok, cell in bindings.items():
                    if cell == (sheet, ref) or tok in claimants:
                        continue
                    b, a = base.raw.get(tok), after.raw.get(tok)
                    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                        if abs(float(a) - float(b)) > 0.005:
                            failures.append(
                                f"{tok} claims {cell[0]}!{cell[1]} but moved when "
                                f"{sheet}!{ref} changed - its source is wrong"
                            )

            assert not failures, (
                f"{len(failures)} binding(s) do not behave as their source claims:\n  "
                + "\n  ".join(failures[:20])
            )
    finally:
        wb.close()
