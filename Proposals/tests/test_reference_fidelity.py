"""Acceptance: the pipeline reproduces the reference proposal's figures.

What this asserts, and what it deliberately does not
----------------------------------------------------
A literal text diff against the reference is the wrong test. Several paragraphs
are now generated from the workbook rather than copied, so their wording differs
on purpose - and it differs in the direction of being correct for whatever site
is being quoted, instead of asserting five Level 2 ports at every one.

What must not differ is the **arithmetic**. So this walks every currency figure
printed in the reference document and requires the pipeline to produce it from
the same workbook. If a number in the original cannot be reproduced, either the
extraction is wrong or the number was never sourced from the workbook - and both
are worth failing over.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.extract import build_context

from .conftest import FOOD4LESS_OPERATOR_INPUTS, REFERENCE_DOCX

# 1 to 4 decimals, not exactly 2. Retail rates print to three ($0.519), and a
# two-decimal pattern silently truncates them to $0.51, which then looks like a
# figure the pipeline failed to reproduce.
CURRENCY = re.compile(r"-?\$[\d,]+(?:\.\d{1,4})?")

# Figures in the reference the pipeline is not expected to reproduce, each with
# the reason. Anything not listed here must reproduce exactly.
KNOWN_ABSENT = {
    # --- idle-fee model ---------------------------------------------------
    # Only the legacy engine has one. The reference prints these in Section 5's
    # KPI strip and its comparison table. The proposal no longer claims idle-fee
    # revenue at all - not in a caption and not in a figure. `idle_fee_annual` is
    # still extracted, because it is a real workbook number and belongs in the QA
    # report, but nothing printed consumes it: the tiles are charger revenue and
    # carbon credits, read off the Financial Worksheet.
    "$26,870", "$108,477", "$564,696", "$433,515",
    "$13,752",   # the comparison table's "Difference" column, idle-inclusive

    # --- hand-rounded prose ------------------------------------------------
    # An analyst wrote "roughly $62,000 in interest" over the real $61,970.48,
    # and similar for the others. The precise figures ARE reproduced; only
    # these rounded restatements are not, and inventing a rounding rule to
    # match them would be fitting the test to one document.
    "$38,000",   # "roughly $38,000 in net profit" over the window
    "$62,000",   # "roughly $62,000 in interest" vs the exact $61,970.48
    "$56,900",   # "roughly $56,900" cumulative net vs the exact cf_net_total
    "$0.30",     # "$0.30 of gross margin", i.e. $0.70 retail less $0.40 cost

    # --- boilerplate constants ---------------------------------------------
    # These describe the market or the product, not this deal.
    "$8,600", "$4,300", "$20", "$35", "$0.0045",
}


def _figures(path: Path) -> set[str]:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    return {m.group(0) for m in CURRENCY.finditer(text)}


@pytest.fixture(scope="module")
def legacy_figures(field_map, legacy_path) -> set[str]:
    """Every currency string the pipeline can produce from the legacy workbook."""
    with engine_mod.load(legacy_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=FOOD4LESS_OPERATOR_INPUTS)

    produced = {str(v) for v in ctx.formatted.values() if str(v).startswith(("$", "-$"))}
    for rows in ctx.tables.values():
        for row in rows:
            for value in row.values():
                if isinstance(value, str) and value.startswith(("$", "-$")):
                    produced.add(value)
    # The document rounds some figures to whole dollars and prints others to the
    # cent, so both spellings of each value count as reproduced.
    for token, raw in list(ctx.raw.items()):
        if isinstance(raw, (int, float)) and not isinstance(raw, bool) and raw:
            sign = "-" if raw < 0 else ""
            produced.add(f"{sign}${abs(raw):,.0f}")
            produced.add(f"{sign}${abs(raw):,.2f}")
    return produced


@pytest.mark.skipif(not REFERENCE_DOCX.exists(), reason="reference proposal missing")
def test_reference_currency_figures_are_reproduced(legacy_figures):
    reference = _figures(REFERENCE_DOCX)
    missing = sorted(f for f in reference - legacy_figures
                     if f not in KNOWN_ABSENT and f not in {"$0", "$0.00"})

    if missing:
        print(f"\n{len(missing)} reference figures not reproduced:")
        for figure in missing:
            print(f"  {figure}")
    assert not missing, (
        f"{len(missing)} figures in the reference proposal cannot be produced "
        "from its own workbook: " + ", ".join(missing[:15])
    )


@pytest.mark.skipif(not REFERENCE_DOCX.exists(), reason="reference proposal missing")
def test_headline_figures_match_exactly(legacy_context):
    """The numbers a reader actually checks."""
    ctx = legacy_context
    assert ctx.formatted["cost_after_discount"] == "$272,023.03"
    assert ctx.formatted["cost_grand_total"] == "$315,180.43"
    assert ctx.formatted["roi_carbon_credits"] == "$218,459.66"
    assert ctx.formatted["roi_itc"] == "$81,606.91"
    assert ctx.formatted["roi_evse_revenues"] == "$406,173.10"
    assert ctx.formatted["roi_net_revenues"] == "$434,216.64"
    assert ctx.formatted["loan_monthly_payment"] == "$5,566.56"
    assert ctx.formatted["loan_total_interest"] == "$61,970.48"
    assert ctx.raw["breakeven_year"] == 4


@pytest.mark.skipif(not REFERENCE_DOCX.exists(), reason="reference proposal missing")
def test_section_3b_matches_the_printed_table(legacy_context):
    """Section 3B is recomputed, so this is the check that it is faithful."""
    rows = {r["label"]: r for r in legacy_context.tables["historical_baseline"]}
    assert rows["Stall Occupancy % (stalls seeing use each day)"]["l2"] == "15.3%"
    assert rows["Stall Occupancy % (stalls seeing use each day)"]["l3"] == "43.0%"
    assert rows["Charger Rating (kW, incl. de-rate)"]["l2"] == "7.06"
    assert rows["Charger Rating (kW, incl. de-rate)"]["l3"] == "49.00"
    assert rows["Total Profit per Year"]["l2"] == "$269.50"
    assert rows["Total Profit per Year"]["l3"] == "$12,848.62"
    assert legacy_context.formatted["hist_profit_yearly"] == "$13,118"
    assert legacy_context.formatted["hist_gross_yearly"] == "$31,220"


@pytest.mark.skipif(not REFERENCE_DOCX.exists(), reason="reference proposal missing")
def test_known_defects_are_not_reproduced(legacy_context):
    """Three faults in the original, deliberately not carried forward.

    1. The cover says August 4 and Section 2 says August 3. One token now.
    2. Every page reads "Page 1" from a stale cached field result.
    3. Section 3B's total block merges the Level 2 and Level 3 columns.
    """
    import zipfile

    template = Path(__file__).resolve().parent.parent / "templates" / "proposal_template.docx"
    if not template.exists():
        pytest.skip("run `python -m tools.build_template` first")

    with zipfile.ZipFile(template) as z:
        document = z.read("word/document.xml").decode("utf-8")
        footer = z.read("word/footer1.xml").decode("utf-8")

    # 1: one date token drives both places
    assert document.count("proposal_date_upper") == 1
    assert "August 3, 2026" not in document and "AUGUST 4, 2026" not in document
    # 2: the cached page number is gone
    assert "1" not in re.findall(r"<w:t[^>]*>([^<]*)</w:t>", footer)
    # 3: 3B renders from a loop with both columns populated per row
    rows = legacy_context.tables["historical_baseline"]
    assert all(r["l2"] and r["l3"] for r in rows)
