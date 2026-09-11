"""Wording that has to change with the site, and captions that must not lie.

Numbers are not the only thing that leaks. "4 single-port 32A units" printed on
a site installing two dual-port 80 A chargers was a sentence with nothing right
in it, and no figure in it for a differential test to catch: `n_ports_l2` was a
PORT count printed with the noun "unit", and `single-port` and `32A` were the
reference site's hardware written as constants.

The other half is captions. A tile reading "Charging profit + idle fees" over a
value containing no idle term is a lie told in the label rather than the
number, and it has happened twice.
"""

from __future__ import annotations

import re

import pytest

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.extract import build_context

from .conftest import (BEST_WESTERN, LEGACY, MARRIOTT, MARRIOTT_L2, SHOWCASE, V16)

BASE = {"site_name": "Site", "client_contact_name": "Contact",
        "prepared_by_name": "Preparer", "site_location_narrative": "N.",
        "existing_ports_l2": 5, "existing_ports_l3": 2}

SPECIMENS = [
    pytest.param(BEST_WESTERN, id="v16"),
    pytest.param(V16, id="v16-master"),
    pytest.param(SHOWCASE, id="v16-legacy-engine"),
    pytest.param(MARRIOTT, id="v15_no_itc"),
    pytest.param(MARRIOTT_L2, id="v16-l2-only"),
    pytest.param(LEGACY, id="legacy_food4less"),
]


def _ctx(path, field_map):
    with engine_mod.load(path, field_map) as lw:
        return build_context(lw, field_map, operator_inputs=dict(BASE))


@pytest.mark.parametrize("path", SPECIMENS)
def test_no_kpi_caption_mentions_idle_fees(path, field_map):
    """Idle fees are out of the projection tiles entirely.

    A caption is read as a definition of the number above it. One that names a
    revenue stream the value does not contain misstates what the figure is,
    which is worse than a wrong figure because it is not checkable.
    """
    ctx = _ctx(path, field_map)
    offenders = {
        t: v for t, v in ctx.formatted.items()
        if t.startswith("kpi_") and isinstance(v, str) and "idle" in v.lower()
    }
    assert not offenders, f"KPI captions still promise idle fees: {offenders}"


@pytest.mark.parametrize("path", SPECIMENS)
def test_itc_is_named_in_a_caption_only_when_there_is_an_itc(path, field_map):
    """`has_itc` gates the value; it must gate the wording too."""
    ctx = _ctx(path, field_map)
    has_itc = bool(ctx.raw.get("has_itc"))
    naming = {
        t: v for t, v in ctx.formatted.items()
        if t.startswith("kpi_") and isinstance(v, str) and "itc" in v.lower()
    }
    if has_itc:
        return
    assert not naming, (
        f"this workbook has no ITC (roi_itc == 0) but a caption still claims "
        f"one: {naming}"
    )


@pytest.mark.parametrize("path", SPECIMENS)
def test_l2_amperage_is_never_invented(path, field_map):
    """Blank is honest; the reference site's 32A is not.

    Printing a current rating on a charger spec sheet that the line items do
    not support is the whole defect this token exists to prevent.
    """
    ctx = _ctx(path, field_map)
    amps = ctx.raw.get("l2_amperage")
    if not amps:
        return  # blank is the correct answer when nothing states it
    items = ctx.tables.get("equipment_line_items") or []
    haystack = " ".join(
        f"{i.get('description') or ''} {i.get('sku') or ''}" for i in items
    )
    digits = re.sub(r"[^0-9]", "", str(amps))
    assert digits and digits in re.sub(r"[^0-9A-Za-z ]", "", haystack), (
        f"l2_amperage is {amps!r} but no line item description or SKU on this "
        f"workbook carries that rating"
    )


@pytest.mark.parametrize("path", SPECIMENS)
def test_l2_sentence_counts_units_not_ports(path, field_map):
    """"4 single-port 32A units" counted ports and called them units."""
    ctx = _ctx(path, field_map)
    sentence = ctx.formatted.get("l2_mix_sentence") or ""
    if not sentence or not ctx.raw.get("has_l2"):
        return
    units = int(ctx.num("n_chargers_l2") or 0)
    if units:
        leading = re.match(r"\s*(\d+)\s*x", sentence)
        assert leading, f"l2_mix_sentence does not lead with a unit count: {sentence!r}"
        assert int(leading.group(1)) == units, (
            f"l2_mix_sentence says {leading.group(1)} but the workbook has "
            f"{units} Level 2 chargers ({ctx.num('n_ports_l2'):.0f} ports)"
        )


@pytest.mark.parametrize("path", SPECIMENS)
def test_carbon_sentence_states_this_workbooks_own_class(path, field_map):
    """Section 13 quoted the reference site's rate and class on every proposal.

    The sentence was frozen literal text - "$8,600 per year for the 100 kW
    class" - while `cc_tier_sentence` computed the right answer and was
    discarded. Both halves have to come from this workbook.
    """
    ctx = _ctx(path, field_map)
    sentence = ctx.formatted.get("cc_tier_sentence")
    assert sentence, "cc_tier_sentence is not published on this chassis"
    if sentence == "the installed classes":
        return  # no L3 tiers on this site; the fallback is honest
    pairs = re.findall(r"\$([\d,]+) per year for the ([\d.]+) kW class", sentence)
    assert pairs, f"cc_tier_sentence has an unexpected shape: {sentence!r}"

    rate = ctx.num("cc_rate_per_kw") or 0
    if rate:
        # Six-tier chassis: every multiplier is rating x the published
        # $/kW/yr, so the sentence must be internally consistent with it.
        for money, kw in pairs:
            implied = float(money.replace(",", "")) / float(kw)
            assert abs(implied - rate) < 0.5, (
                f"{sentence!r} implies ${implied:.2f}/kW/yr but the workbook's "
                f"cc_rate_per_kw is ${rate}"
            )
        return

    # Single-tier chassis (legacy): there is no rate grid at all - the tier
    # quantities and rates are hardcoded inside the B4 formula, and the
    # multiplier does NOT follow $/kW (Food4Less carries $8,600 against a 98 kW
    # charger, which is $87.76/kW). So check the sentence against the two
    # scalars it was built from instead.
    money, kw = pairs[0]
    assert abs(float(money.replace(",", "")) - ctx.num("cc_l3_mult")) < 1.0, (
        f"{sentence!r} does not carry this workbook's cc_l3_mult "
        f"({ctx.num('cc_l3_mult')})")
    assert abs(float(kw) - ctx.num("cc_l3_rating")) < 0.5, (
        f"{sentence!r} does not carry this workbook's cc_l3_rating "
        f"({ctx.num('cc_l3_rating')})")


# ---------------------------------------------------------------------------
# Idle fees are not part of this proposal at all.
#
# They were removed in stages: first from the projection tiles' arithmetic,
# then from their captions, and finally - here - from Section 4's methodology
# and assumptions tables and Section 10's EVOLV capability list. A document
# that describes an idle-fee policy while modelling no idle-fee revenue is
# describing a revenue stream it does not claim, and on a scenario-grid
# workbook the grace-period row rendered as "0-minute Level 2 / 0-minute
# Level 3 grace", which is worse than saying nothing.
# ---------------------------------------------------------------------------

IDLE_WORDS = re.compile(r"idle|grace[- ]period|overstay", re.I)


@pytest.mark.parametrize("path", SPECIMENS)
def test_no_rendered_proposal_mentions_idle_fees(path, field_map, tmp_path):
    from docx import Document

    from ev_proposal_agent.render import generate

    out = tmp_path / "proposal.docx"
    generate(path, output=out, force=True, progress=lambda _m: None,
             operator_inputs=dict(BASE))

    doc = Document(out)
    blob: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            blob.extend(c.text for c in row.cells)
    for section in doc.sections:
        for part in (section.header, section.footer):
            blob.extend(p.text for p in part.paragraphs)

    hits = sorted({" ".join(t.split())[:90] for t in blob if IDLE_WORDS.search(t)})
    assert not hits, "the proposal still describes an idle-fee policy:\n  " + "\n  ".join(hits)


def test_the_methodology_table_numbering_has_no_gap(field_map, tmp_path):
    """Deleting layer 5 must renumber 6 and 7, not leave 1,2,3,4,6,7."""
    from docx import Document

    from ev_proposal_agent.render import generate

    out = tmp_path / "proposal.docx"
    generate(BEST_WESTERN, output=out, force=True, progress=lambda _m: None,
             operator_inputs=dict(BASE))

    doc = Document(out)
    table = next((t for t in doc.tables
                  if " ".join(t.rows[0].cells[0].text.split()) == "Layer"), None)
    assert table is not None, "Section 4's methodology table is not in the document"

    ordinals = []
    for row in table.rows:
        m = re.match(r"^(\d+)\.", " ".join(row.cells[0].text.split()))
        if m:
            ordinals.append(int(m.group(1)))
    assert ordinals == list(range(1, len(ordinals) + 1)), (
        f"the methodology layers are numbered {ordinals}, which has a gap")


# ---------------------------------------------------------------------------
# A continuation heading names its parent in prose - "Continuation of Section
# 7" - and that number was a frozen literal sitting beside a badge that was a
# token. On a greenfield workbook, Sections 3/3A/3B are suppressed and every
# later section moves up, so 6A read "Continuation of Section 7" and 16B read
# "Appendix to Section 17". Seven of nine were wrong on the first real
# greenfield file, and nothing caught it: the badge, the contents listing and
# the Section 2 grid all came from one plan, and this was a fourth place that
# did not.
# ---------------------------------------------------------------------------

CONTINUATION = re.compile(r"(?:Continuation of|Appendix to) Section (\d+)")


@pytest.mark.parametrize("path", SPECIMENS)
def test_continuation_headings_name_their_real_parent(path, field_map, tmp_path):
    from docx import Document

    from ev_proposal_agent.render import generate

    out = tmp_path / "proposal.docx"
    generate(path, output=out, force=True, progress=lambda _m: None,
             operator_inputs=dict(BASE))

    doc = Document(out)
    wrong, checked = [], 0
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.rows[0].cells) != 2:
            continue
        badge = table.rows[0].cells[0].text.strip()
        title = " ".join(table.rows[0].cells[1].text.split())
        if not badge or not title:
            continue
        m = CONTINUATION.search(title)
        if not m:
            continue
        checked += 1
        parent = re.sub(r"[A-Z]$", "", badge)
        if m.group(1) != parent:
            wrong.append(f"{badge} says {m.group(0)!r} but its parent is {parent}")

    assert checked >= 4, (
        f"only {checked} continuation headings found - the test has stopped "
        f"covering anything")
    assert not wrong, "continuation headings cite the wrong section:\n  " + "\n  ".join(wrong)
