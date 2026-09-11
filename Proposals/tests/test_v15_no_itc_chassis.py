"""The `v15_no_itc` chassis - Marriott Bakersfield.

A hybrid, and the reason `is_legacy_chassis` stopped deciding anything: it takes
the legacy cost block and cashflow rows AND the v16 INPUT SHEET and EVOLV block
at the same time, which no boolean can express.

Every test below names the specific way one of the other two chassis would get
this file wrong, because "it produced a number" is not the bar - both wrong
chassis produce numbers here, and they are plausible ones.
"""

from __future__ import annotations

import openpyxl
import pytest

from ev_proposal_agent.validate import validate

from .conftest import MARRIOTT


def _codes(findings) -> set[str]:
    return {f.code for f in findings}


def test_detected_as_its_own_chassis(marriott_context):
    det = marriott_context.detection
    assert det.chassis == "v15_no_itc"
    # It is NOT the legacy chassis, even though its cost rows are legacy's.
    assert det.is_legacy_chassis is False
    # The revenue model is a separate axis and is the legacy one.
    assert det.engine == "baseline_vs_projected"


def test_no_itc_row_is_not_read_as_evse_revenues(marriott_context):
    """The regression this chassis exists to prevent.

    There is no ITC row. Under the legacy chassis, resolve() tries the label
    'ITC', misses, falls through to the fixed cell B5 - which here holds EVSE
    Revenues - and Section 7 prints $416,761 as a Federal ITC line with nothing
    raised.
    """
    ctx = marriott_context
    assert ctx.num("roi_itc") == 0.0
    assert ctx.raw["has_itc"] is False
    assert ctx.raw["has_itc_row"] is False
    assert "absent" in ctx.sources["roi_itc"]
    assert ctx.num("roi_itc") != pytest.approx(416761, abs=1)


def test_evse_and_net_come_from_b5_b6_by_address(marriott_context):
    """One row up from every other chassis, and addressed, not label-searched."""
    ctx = marriott_context
    assert ctx.sources["roi_evse_revenues"] == "Financial Worksheet!B5"
    assert ctx.sources["roi_net_revenues"] == "Financial Worksheet!B6"
    # B6 = B3+(B4+B5). No ITC term to include or exclude.
    assert ctx.num("roi_net_revenues") == pytest.approx(
        ctx.num("roi_total_costs_upfront")
        + ctx.num("roi_carbon_credits")
        + ctx.num("roi_evse_revenues"),
        abs=0.01,
    )


def test_cost_block_uses_the_legacy_rows(marriott_context):
    """v16 would read B43/B44 and find the discount total and a blank."""
    ctx = marriott_context
    assert ctx.sources["cost_grand_total"] == "Financial Worksheet!B42"
    assert ctx.sources["cost_after_discount"] == "Financial Worksheet!B43"


def test_cost_hierarchy_foots_to_after_discount(marriott_context, field_map):
    """The components are customer prices, so they foot to the DISCOUNTED total.

    Asserting the two totals differ is what makes this test mean something: this
    is a discounted workbook, so rooting the check at the grand total (as it was)
    fails here by $91,531.
    """
    ctx = marriott_context
    assert "cost-subtotal-mismatch" not in _codes(validate(ctx, field_map))

    components = sum(
        ctx.num(t) for t in ("cost_equipment_invoice", "cost_design_invoice",
                             "cost_electrical_subtotal", "cost_labor")
    )
    assert ctx.num("cost_after_discount") == pytest.approx(components, abs=1)
    assert ctx.num("cost_grand_total") != pytest.approx(components, abs=1)


def test_cashflow_uses_the_legacy_rows(marriott_context):
    """v16's header is row 4 and its data starts at 5; this file starts at 4."""
    rows = marriott_context.tables["consolidated_cashflow"]
    assert rows[0]["_row"] == 4
    assert len(rows) == 11          # year 0 through year 10


def test_evolv_uses_the_v16_block(marriott_context, field_map):
    """The legacy override reads G4 - the '# of years' cell - and quotes 5 ports
    for a 14-port site, blocking with evolv-port-mismatch."""
    ctx = marriott_context
    assert ctx.sources["evolv_ports"] == "Internal Summary!G3"
    assert ctx.num("evolv_ports") == 14
    assert ctx.num("evolv_ports") == ctx.num("total_ports")
    assert ctx.raw["evolv_fee_per_port"] is not None
    assert ctx.raw["evolv_annual_cost"] is not None
    assert "evolv-port-mismatch" not in _codes(validate(ctx, field_map))


def test_service_term_comes_from_g4_not_g5(marriott_context):
    """No F9:I14 service-plan block, as on legacy - but the EVOLV block above it
    is the v16 one, so the contract length sits at G4, not legacy's G5."""
    ctx = marriott_context
    assert ctx.sources["svc_contract_years"] == "Internal Summary!G4"
    assert ctx.num("svc_contract_years") == 5


def test_carbon_grid_is_six_tier_and_relocated(marriott_context):
    """v16 structure at new addresses. Both existing maps read empty cells here:
    v16's D42:J47 and legacy's I43:K46 are blank on this file, which left the
    Section 13 tile reading '$0'."""
    ctx = marriott_context
    assert ctx.sources["cc_tier_mult"] == "Financial Worksheet!M7:R7"
    assert ctx.sources["cc_tier_qty"] == "Financial Worksheet!M6:R6"
    for token in ("cc_rate_per_kw", "cc_l2_kwh_month", "cc_l2_rate",
                  "cc_l2_credit_month"):
        assert isinstance(ctx.raw[token], (int, float)), token

    assert ctx.raw["cc_multiplier_lines"] != "$0"
    assert ctx.raw["cc_tier_sentence"]
    tiers = ctx.raw["_carbon_tiers"]
    assert len(tiers) == 1                      # only the 60 kW class is bought
    assert tiers[0]["kw"] == 60
    assert tiers[0]["qty"] == 4


def test_carbon_total_reconciles_with_the_grid(marriott_context, marriott_path):
    """Proves the relocated addresses are the RIGHT ones, not merely readable.

    FW!B4 = (10*SUMPRODUCT(M6:R6,M7:R7)) + (O9*'Updated'!D6*10*12) - the v16
    carbon formula exactly, same terms, new addresses.
    """
    ctx = marriott_context
    wb = openpyxl.load_workbook(marriott_path, data_only=True)
    projected_l2_ports = wb["Updated Chargers Revenue Calcul"]["D6"].value
    wb.close()

    years = ctx.num("projection_years")
    l3 = sum(q * m for q, m in zip(ctx.raw["cc_tier_qty"], ctx.raw["cc_tier_mult"]))
    l2 = ctx.num("cc_l2_credit_month") * projected_l2_ports * years * 12

    assert ctx.num("roi_carbon_credits") == pytest.approx(years * l3 + l2, abs=1)


def test_horizon_is_fixed_ten_and_the_divisor_defect_is_reported(marriott_context):
    """Fixed 10-year, and `Cashflow!G3` still divides by `(5*12)`.

    Having no B9/B10 toggle does not make the divisor right - it just means the
    file was always wrong rather than becoming wrong when someone flipped a
    switch. B4 holds ten years of carbon; G3 spreads it over five, so the
    60-month tables carry $2,975.83/month against a true $1,487.91.

    Fix `G3` to `='Financial Worksheet'!$B$4/(10*12)` in the workbook - or add
    the B9/B10 toggle and use `$B$10*12` - and this goes quiet.
    """
    ctx = marriott_context
    assert ctx.detection.projection_years == 10
    assert ctx.detection.has_financing is True
    assert "cashflow-carbon-doubled" in _codes(ctx.findings)


def test_utility_rate_is_read_not_assumed(marriott_context):
    """This workbook's EV utility rate is $0.2584/kWh, not the $0.40 the other
    two use. Assuming 0.40 gave $10,979.55 while Section 5 and the ten-year
    headline in the same document said $16,446.75."""
    ctx = marriott_context
    assert ctx.raw["hist_utility_rate_l2"] == pytest.approx(0.2584, abs=1e-6)
    assert ctx.num("hist_profit_yearly") == pytest.approx(16446.75, abs=0.5)
    assert ctx.num("hist_profit_yearly") != pytest.approx(10979.55, abs=1)


def test_recompute_matches_this_workbooks_updated_sheet(marriott_context, marriott_path):
    """Section 3B is recomputed in Python; it must land on the workbook's own
    Updated!B20 + C20 or the proposal contradicts itself."""
    wb = openpyxl.load_workbook(marriott_path, data_only=True)
    ws = wb["Updated Chargers Revenue Calcul"]
    expected = ws["B20"].value + ws["C20"].value
    wb.close()
    assert marriott_context.num("hist_profit_yearly") == pytest.approx(expected, abs=0.01)


def test_existing_ports_are_prefilled_from_the_updated_sheet(marriott_context):
    """No operator input was supplied, so these came off Updated!B6/C6 - the
    EXISTING stall counts, not Internal Summary!G3's 14 new chargers."""
    ctx = marriott_context
    assert ctx.num("existing_ports_l2") == 12
    assert ctx.num("existing_ports_l3") == 4


def test_amortization_is_rebuilt_from_the_summary_inputs(marriott_context, marriott_path):
    """This workbook's DLL amortization table sits one column left of its own
    headers, so `Total_Interest` summed the ending-balance column and read
    $10,904,157.52 on a $253,717 loan. The schedule is recomputed from the
    D5:D10 summary block instead.

    The rebuild is only trusted because it reproduces two figures the broken
    table had no hand in - the workbook's own `J5` payment and `Cashflow!F1`.
    """
    ctx = marriott_context
    rows = ctx.tables["amortization"]

    wb = openpyxl.load_workbook(marriott_path, data_only=True)
    stated_payment = wb["DLL Schedule"]["J5"].value
    cashflow_total = wb["Cashflow"]["F1"].value
    wb.close()

    assert len(rows) == 60
    assert rows[0]["begin"] == pytest.approx(ctx.num("loan_amount"), abs=0.01)
    assert rows[-1]["end"] == pytest.approx(0.0, abs=0.01)
    assert rows[0]["sched"] == pytest.approx(stated_payment, abs=0.005)
    assert sum(r["total"] for r in rows) == pytest.approx(cashflow_total, abs=0.05)

    assert ctx.num("loan_total_interest") == pytest.approx(57800.09, abs=0.05)
    assert ctx.num("loan_total_interest") < ctx.num("loan_amount")
    assert ctx.num("loan_actual_payments") == 60
    assert ctx.num("loan_early_payments") == 0
    assert "amortization-rebuilt" in _codes(ctx.findings)


def test_financing_sections_survive_the_rebuild(marriott_context):
    """Repaired, not dropped - all 28 sections render."""
    assert marriott_context.raw["has_financing"] is True


def test_no_error_findings(marriott_context, field_map):
    """The acceptance gate, minus the one defect the workbook itself has.

    `cashflow-carbon-doubled` is expected here and is the operator's to fix
    in Excel (see the divisor test above). Everything else must be clean.
    """
    errors = [f for f in validate(marriott_context, field_map)
              if f.level == "ERROR" and f.code != "cashflow-carbon-doubled"]
    assert errors == [], [f"{f.code}: {f.message}" for f in errors]


def test_workbook_is_in_the_corpus():
    assert MARRIOTT.exists()


# --------------------------------------------------------------------------
# The same chassis carrying the OTHER engine, and a moved carbon block.
#
# Rip & Replace_MarriotbakersfieldL2 is v15_no_itc with the scenario grid, no
# Historical Data, no 10 Year Projection Comparison, and its carbon grid one row
# below the other specimen's. It is the file that proved two things: that the
# chassis cannot require either optional sheet, and that a fixed carbon address
# reads the neighbouring row on whichever file it was not written for.
# --------------------------------------------------------------------------


def test_l2_specimen_is_the_same_chassis_with_the_scenario_grid(marriott_l2_context):
    det = marriott_l2_context.detection
    assert det.chassis == "v15_no_itc"
    assert det.engine == "scenario_grid"
    assert det.has_history is False
    assert det.ten_year_sheet_present is False


def test_l2_scenario_is_read_from_i5_not_i6(marriott_l2_context):
    """I6 holds `=I5*1.125` on this chassis - a growth formula with no reference
    to the Updated tab. Reading it warns and falls back to Standard-Low, which
    would be silently wrong on a workbook wired to Medium or High."""
    ctx = marriott_l2_context
    assert ctx.detection.scenario == "Standard-Low"
    assert "scenario-unreadable" not in _codes(ctx.findings)
    assert "scenario-unparsed" not in _codes(ctx.findings)


def test_l2_carbon_block_is_found_one_row_lower(marriott_l2_context):
    """L5:R10 here against L4:R9 on the historical specimen. Same offsets, found
    from the label, so neither file needs its own address map."""
    ctx = marriott_l2_context
    assert ctx.sources["cc_rate_per_kw"] == "Financial Worksheet!M5"
    assert ctx.sources["cc_tier_qty"] == "Financial Worksheet!M7:R7"
    assert ctx.sources["cc_l2_credit_month"] == "Financial Worksheet!O10"
    assert "carbon-block-not-found" not in _codes(ctx.findings)


def test_l2_carbon_total_reconciles(marriott_l2_context, marriott_l2_path):
    """This site buys no Level 3 at all, so the whole credit is the L2 term.
    Getting the block one row out would read the ratings row as quantities and
    invent six tiers of DC fast credits for a single Level 2 charger."""
    ctx = marriott_l2_context
    wb = openpyxl.load_workbook(marriott_l2_path, data_only=True)
    l2_stalls = wb["Updated Chargers Revenue Calcul"]["B6"].value
    wb.close()

    assert sum(ctx.raw["cc_tier_qty"]) == 0
    assert ctx.raw["_carbon_tiers"] == []
    years = ctx.num("projection_years")
    expected = ctx.num("cc_l2_credit_month") * l2_stalls * years * 12
    assert ctx.num("roi_carbon_credits") == pytest.approx(expected, abs=0.01)


def test_l2_gross_revenue_is_derived_not_blank(marriott_l2_context):
    """The scenario grid's totals block has no gross-revenue row, so Section 5's
    last line rendered empty. It is 'Per 30 day Cycle' x 12, both levels."""
    ctx = marriott_l2_context
    rows = {r["token"]: r for r in ctx.tables["operating_model"]}
    per_30 = rows["revenue_per_30day"]
    assert ctx.num("projected_gross_yearly") == pytest.approx(
        (per_30["l2"] + per_30["l3"]) * 12, abs=0.01)


def test_l2_has_no_error_findings(marriott_l2_context, field_map):
    """The acceptance gate, minus the one defect the workbook itself has.

    `cashflow-carbon-doubled` is expected here and is the operator's to fix
    in Excel (see the divisor test above). Everything else must be clean.
    """
    errors = [f for f in validate(marriott_l2_context, field_map)
              if f.level == "ERROR" and f.code != "cashflow-carbon-doubled"]
    assert errors == [], [f"{f.code}: {f.message}" for f in errors]