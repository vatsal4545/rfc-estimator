"""The real production shape: a v16 chassis carrying the legacy revenue engine.

Showcase Liquor Pasadena is what operators will actually drop in. It matters
because it breaks the assumption the first design was built on - that the
workbook version and the revenue model are the same fact. They are not:

    chassis  v16                    cost block B16-B44, cashflow rows 5-15,
                                    B9/B10 horizon toggle, EVOLV at G3,
                                    service-plan block present, carbon D42:J47
    engine   baseline_vs_projected  Updated holds B/C historical, D/E projected,
                                    F:H idle fees - not the seven-column grid

It also covers two data shapes the other two files do not: a site with no DC
fast history whatsoever, and a station-count note that mentions only Level 2.
"""

from __future__ import annotations

import openpyxl
import pytest

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.extract import build_context

from .conftest import SHOWCASE_OPERATOR_INPUTS

CENTS = 0.005


# --------------------------------------------------------------------------
# the two axes
# --------------------------------------------------------------------------


def test_v16_chassis_with_legacy_engine(showcase_context):
    """The combination the first design rejected outright."""
    det = showcase_context.detection
    assert det.chassis == "v16"
    assert det.engine == "baseline_vs_projected"
    assert det.is_legacy_engine is True
    assert det.is_legacy_chassis is False


def test_chassis_row_layout_is_v16(showcase_context):
    """Cost block one row LOWER than the legacy chassis. Reading legacy
    addresses here would return Labor as the grand total."""
    ctx = showcase_context
    assert ctx.sources["cost_grand_total"] == "Financial Worksheet!B43"
    assert ctx.sources["cost_after_discount"] == "Financial Worksheet!B44"
    assert ctx.num("cost_grand_total") == pytest.approx(49988.10125, abs=CENTS)
    assert ctx.num("cost_after_discount") == pytest.approx(49988.10125, abs=CENTS)


def test_cashflow_uses_v16_rows(showcase_context):
    """Header row 4, data from row 5 - not the legacy row 3 / row 4."""
    rows = showcase_context.tables["consolidated_cashflow"]
    assert rows[0]["_row"] == 5
    assert len(rows) == showcase_context.raw["projection_years"] + 1 == 6


def test_horizon_toggle_is_honoured(showcase_context):
    """The legacy chassis has no B9/B10 toggle and is fixed at 10 years. This
    file is a v16 chassis set to 5, and must be read as 5."""
    assert showcase_context.raw["projection_years"] == 5


def test_v16_only_blocks_are_read(showcase_context):
    """EVOLV at G3 (not G4) and the service-plan block, neither of which the
    legacy chassis has."""
    ctx = showcase_context
    assert ctx.sources["evolv_ports"] == "Internal Summary!G3"
    assert ctx.num("evolv_ports") == 2
    assert ctx.num("svc_contract_years") == 5
    assert ctx.num("cost_service_warranty") == pytest.approx(13118.87, abs=CENTS)


def test_no_scenario_is_detected_for_the_legacy_engine(showcase_context):
    """`Financial Worksheet!I6` points at `Updated!C24`, the legacy projected
    total. Trying to read that as a scenario grid is exactly the collision."""
    assert showcase_context.detection.scenario is None


def test_operating_model_uses_the_projected_columns(showcase_context):
    """Legacy engine: columns D and E are already what Section 5 prints."""
    rows = showcase_context.tables["operating_model"]
    tokens = [r["token"] for r in rows]
    assert "avg_delivered_pct" in tokens, "legacy row 13 exists on this sheet"
    by_token = {r["token"]: r for r in rows}
    assert by_token["chargers"]["l2"] == 1
    assert by_token["chargers"]["l3"] == 1
    assert by_token["profit_per_year"]["l3"] == pytest.approx(6350.40, abs=CENTS)


# --------------------------------------------------------------------------
# the ten-year sheet is CORRECT here
# --------------------------------------------------------------------------


def test_ten_year_sheet_is_readable_not_blocked(showcase_context):
    """Its formulas point at the legacy Updated addresses, and the legacy
    Updated sheet is what this workbook has. It is only mis-wired, and only
    refused, when the scenario grid is present instead."""
    det = showcase_context.detection
    assert det.ten_year_sheet_present is True
    assert det.ten_year_sheet_readable is True
    codes = {f.code for f in showcase_context.findings}
    assert "ten-year-sheet-miswired" not in codes


def test_no_error_findings(showcase_context):
    """A clean production workbook must generate without an override."""
    assert showcase_context.findings.errors == []


# --------------------------------------------------------------------------
# figures cross-checked against the workbook's own cells
# --------------------------------------------------------------------------


def test_recompute_matches_this_workbooks_updated_sheet(showcase_path, showcase_context):
    """Independent confirmation on a second real file: the Python recompute
    reproduces `Updated!B8/B20/B27` without reading them."""
    wb = openpyxl.load_workbook(showcase_path, data_only=True)
    ws = wb["Updated Chargers Revenue Calcul"]
    l2 = showcase_context.raw["_baseline_l2"]

    assert l2.stall_occupancy == pytest.approx(float(ws["B8"].value), rel=1e-9)
    assert l2.charger_rating_kw == pytest.approx(float(ws["B12"].value), rel=1e-9)
    assert l2.net_profit_month == pytest.approx(float(ws["B19"].value), rel=1e-9)
    assert l2.profit_per_year == pytest.approx(float(ws["B20"].value), rel=1e-9)
    assert showcase_context.num("hist_profit_yearly") == pytest.approx(
        float(ws["B24"].value), rel=1e-9)
    assert showcase_context.num("hist_gross_yearly") == pytest.approx(
        float(ws["B27"].value), rel=1e-9)
    wb.close()


@pytest.mark.parametrize(
    "token, expected",
    [
        ("roi_evse_revenues", 42572.20549237338),
        ("roi_net_revenues", 14109.732758286344),
        ("roi_carbon_credits", 21525.628515912962),
        ("roi_itc", 0.0),
        ("year1_charger_profit", 6635.050176810675),
        ("hist_profit_yearly", 774.3280),
    ],
)
def test_headline_figures(showcase_context, token, expected):
    assert showcase_context.num(token) == pytest.approx(expected, abs=0.01)


def test_client_info_prefills_from_the_workbook(showcase_context):
    """Unlike the blank v16, this file has the INPUT SHEET block filled in."""
    ctx = showcase_context
    assert ctx.raw["site_address"] == "1392 N Lake Ave, Pasadena, CA 91104"
    assert ctx.raw["scope_of_work"].startswith("Yanking 3 lvl 2 out")
    codes = {f.code for f in ctx.findings}
    assert "site-info-empty" not in codes


# --------------------------------------------------------------------------
# a site with no DC fast history
# --------------------------------------------------------------------------


def test_history_window_parsed_from_this_files_formula(showcase_context):
    """`ROWS(A4:A38)` here, `ROWS(A8:A42)` in the other specimen. The window is
    per-file and must be read, not assumed."""
    ctx = showcase_context
    assert ctx.raw["months_in_window"] == 35
    assert ctx.raw["history_window_label"] == "Sep 2023 to Jul 2026"
    assert len(ctx.tables["monthly_series"]) == 35


def test_trailing_partial_month_excluded(showcase_context):
    """Row 39 is a partial 2026-08 carrying $55.14, and row 40 is a Total that
    sums everything. Neither belongs in the window."""
    ctx = showcase_context
    assert ctx.num("hist_total_revenue") == pytest.approx(10424.86, abs=0.5)
    assert ctx.num("hist_total_revenue") != pytest.approx(10480.00, abs=0.5)


def test_level_with_no_history_yields_zeroes_not_a_crash(showcase_context):
    """This site never had a DC fast charger. Level 3 is a real baseline of
    nothing, which must not divide by zero or raise."""
    l3 = showcase_context.raw["_baseline_l3"]
    assert l3.profit_per_year == 0.0
    assert l3.revenue_per_day == 0.0
    assert l3.stall_occupancy == 0.0
    assert showcase_context.num("dcfc_revenue_share") == 0.0


def test_baseline_table_still_renders_both_columns(showcase_context):
    rows = showcase_context.tables["historical_baseline"]
    assert len(rows) == 17
    assert all(r["l2"] and r["l3"] for r in rows), "zero must print, not blank"


def test_station_note_with_only_level_2_is_parsed(showcase_context):
    """The note reads "4 Level 2 station IDs" and never mentions Level 3. A
    regex demanding both levels matches nothing and silently skips the check."""
    codes = {f.code for f in showcase_context.findings}
    assert "port-count-disagrees-with-sheet" not in codes


def test_dual_port_stations_do_not_trip_the_port_check(field_map, showcase_path):
    """4 station IDs against 8 ports is two ports per station - normal dual
    connector hardware, not a discrepancy."""
    with engine_mod.load(showcase_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=SHOWCASE_OPERATOR_INPUTS)
    assert "port-count-disagrees-with-sheet" not in {f.code for f in ctx.findings}


def test_a_genuinely_odd_port_count_still_warns(field_map, showcase_path):
    """7 ports across 4 stations is 1.75 each, which is not real hardware."""
    odd = dict(SHOWCASE_OPERATOR_INPUTS, existing_ports_l2=7)
    with engine_mod.load(showcase_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=odd)
    assert "port-count-disagrees-with-sheet" in {f.code for f in ctx.findings}


def test_history_without_ports_is_flagged(field_map, showcase_path):
    """The sheet records Level 2 revenue; entering zero Level 2 ports would
    silently zero out the entire historical baseline."""
    contradictory = dict(SHOWCASE_OPERATOR_INPUTS, existing_ports_l2=0,
                         existing_ports_l3=3)
    with engine_mod.load(showcase_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=contradictory)
    assert "baseline-ports-contradict-history" in {f.code for f in ctx.findings}


# --------------------------------------------------------------------------
# defects carried by this chassis too
# --------------------------------------------------------------------------


def test_broken_service_agreement_cell_does_not_abort(showcase_context):
    """`Financial Worksheet!B20` is a live `#REF!` here as well. It maps to an
    emit:false token, so the "no token starts with #" gate must skip it."""
    ctx = showcase_context
    assert "cost_service_list_price" not in ctx.emitted
    emitted = [v for k, v in ctx.raw.items() if k in ctx.emitted and isinstance(v, str)]
    assert not any(s.strip().startswith("#") for s in emitted)


def test_carbon_divisor_guard_stays_quiet_at_five_years(showcase_context):
    """`Cashflow!G3` divides by (5*12) and the horizon is 5, so the figures are
    right. The guard only fires when the toggle is flipped to 10."""
    assert "cashflow-carbon-doubled" not in {f.code for f in showcase_context.findings}
