"""Phase 1 acceptance: every figure the brief pins down, read from the real file.

If one of these drifts, either the workbook changed or the row map is wrong.
Both are worth failing a build over.
"""

from __future__ import annotations

import pytest

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.errors import FingerprintMismatch, UnknownEngine
from ev_proposal_agent.extract import build_context

CENTS = 0.005


# --------------------------------------------------------------------------
# detection
# --------------------------------------------------------------------------


def test_detects_v16_chassis(v16_context):
    assert v16_context.detection.chassis == "v16"
    assert v16_context.detection.engine == "scenario_grid"


def test_engine_is_scenario_grid(v16_context):
    assert v16_context.raw["engine"] == "scenario_grid"


def test_scenario_is_standard_low(v16_context):
    """Parsed from the FORMULA of Financial Worksheet!I6, not from a value."""
    assert v16_context.raw["revenue_scenario"] == "Standard-Low"


def test_projection_years_is_five(v16_context):
    assert v16_context.raw["projection_years"] == 5


def test_v16_has_no_historical_data_sheet(v16_context):
    """v16 ships without it; Sections 3/3A/3B suppress rather than fail."""
    assert v16_context.raw["has_history"] is False
    codes = {f.code for f in v16_context.findings}
    assert "history-absent" in codes


# --------------------------------------------------------------------------
# costs
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "token, expected",
    [
        ("cost_charger_hardware", 28500.00),
        ("cost_service_warranty", 10270.32),
        ("cost_grand_total", 43235.97),
        ("cost_after_discount", 43235.97),
    ],
)
def test_cost_values(v16_context, token, expected):
    assert v16_context.num(token) == pytest.approx(expected, abs=CENTS)


def test_cost_stack_is_hierarchical_not_a_flat_sum(v16_context):
    """Grand Total is four subtotals, not the sum of every printed row.

    B16 alone already contains B17..B21, so adding the rows up double-counts
    by an entire equipment invoice.
    """
    ctx = v16_context
    hierarchy = ctx.raw["_cost_hierarchy"]
    for parent, children in hierarchy.items():
        expected = sum(ctx.num(child) for child in children)
        assert ctx.num(parent) == pytest.approx(expected, abs=1.0), parent

    flat = sum(row["value"] for row in ctx.raw["_cost_rows"]
               if row["token"] != "cost_after_discount")
    assert flat > ctx.num("cost_grand_total"), "flat sum should over-count; the check must be hierarchical"


def test_service_agreement_prices_match_at_zero_discount(v16_context):
    """B18 is post-discount (IS!D5), B20 is list (IS!B5). Equal at 0% discount."""
    assert v16_context.num("cost_service_warranty") == pytest.approx(
        v16_context.num("cost_service_list_price"), abs=CENTS
    )


# --------------------------------------------------------------------------
# ROI
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "token, expected",
    [
        ("roi_carbon_credits", 21500.01),
        ("roi_itc", 0.00),
        ("roi_evse_revenues", 40745.816),
        ("roi_net_revenues", 19009.856),
    ],
)
def test_roi_values(v16_context, token, expected):
    assert v16_context.num(token) == pytest.approx(expected, abs=CENTS)


def test_itc_row_is_suppressed(v16_context):
    """B5 == 0, so the Federal ITC row and stat card drop out."""
    assert v16_context.num("roi_itc") == 0.0
    assert v16_context.raw["has_itc"] is False


def test_net_revenues_excludes_itc_by_design(v16_context):
    """Financial Worksheet!B7 is =B3+(B4+B6). ITC is deliberately outside it."""
    ctx = v16_context
    expected = ctx.num("roi_total_costs_upfront") + (
        ctx.num("roi_carbon_credits") + ctx.num("roi_evse_revenues")
    )
    assert ctx.num("roi_net_revenues") == pytest.approx(expected, abs=CENTS)


# --------------------------------------------------------------------------
# financing
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "token, expected",
    [
        ("loan_monthly_payment", 884.7617),
        ("loan_total_interest", 9849.7323),
        ("loan_amount", 43235.97),
    ],
)
def test_loan_values(v16_context, token, expected):
    assert v16_context.num(token) == pytest.approx(expected, abs=0.001)


def test_loan_defined_names_are_sheet_scoped(field_map, v16_paths):
    """wb.defined_names['Loan_Amount'] raises KeyError; resolve() must look on
    the DLL Schedule sheet scope. This is the trap that silently breaks the
    `name:` entries in the map."""
    from ev_proposal_agent.resolve import find_defined_name

    with engine_mod.load(v16_paths, field_map) as lw:
        assert "Loan_Amount" not in lw.values.defined_names
        assert find_defined_name(lw.values, "Loan_Amount") is not None
        # cumCash* really are workbook-scoped, unlike the loan names
        assert "cumCashCats" in lw.values.defined_names


# --------------------------------------------------------------------------
# cashflow and equipment
# --------------------------------------------------------------------------


def test_year1_charger_profit(v16_context):
    assert v16_context.num("year1_charger_profit") == pytest.approx(6350.40, abs=CENTS)


def test_breakeven_year(v16_context):
    assert v16_context.raw["breakeven_year"] == 4


def test_consolidated_cashflow_length(v16_context):
    """Year 0 plus the horizon. Rows past it cache as None and are dropped."""
    rows = v16_context.tables["consolidated_cashflow"]
    assert len(rows) == 6
    assert len(rows) == v16_context.raw["projection_years"] + 1


def test_cashflow_cumulative_starts_at_negative_project_cost(v16_context):
    rows = v16_context.tables["consolidated_cashflow"]
    assert rows[0]["cumulative"] == pytest.approx(
        -v16_context.num("cost_after_discount"), abs=1.0
    )


def test_evse_revenues_equals_sum_of_charger_years(v16_context):
    rows = v16_context.tables["charger_cashflow"]
    total = sum(r["annual"] for r in rows if r["year"] and r["year"] >= 1)
    assert total == pytest.approx(v16_context.num("roi_evse_revenues"), abs=1.0)


def test_total_ports(v16_context):
    assert v16_context.num("total_ports") == 1


def test_modeled_rating_l3_is_charger_weighted(v16_context):
    """One active 60 kW tier at a 0.98 de-rate."""
    assert v16_context.num("modeled_rating_l3") == pytest.approx(58.8, abs=0.01)


def test_level_2_column_suppressed_when_no_l2_chargers(v16_context):
    assert v16_context.num("n_ports_l2") == 0
    assert v16_context.raw["has_l2"] is False


def test_evolv_ports_match_total_ports(v16_context):
    assert v16_context.num("evolv_ports") == v16_context.num("total_ports")


def test_evolv_annual_cost_is_consistent(v16_context):
    ctx = v16_context
    expected = ctx.num("evolv_ports") * ctx.num("evolv_fee_per_port") * 12
    assert ctx.num("evolv_annual_cost") == pytest.approx(expected, abs=0.05)


# --------------------------------------------------------------------------
# the legacy path
# --------------------------------------------------------------------------


def test_legacy_workbook_takes_the_legacy_path(legacy_context):
    ctx = legacy_context
    assert ctx.detection.chassis == "legacy_food4less"
    assert ctx.raw["engine"] == "baseline_vs_projected"
    assert ctx.raw["projection_years"] == 10
    assert ctx.raw["has_history"] is True


def test_legacy_cost_rows_are_shifted_up_one(legacy_context):
    """Grand Total is B42 in legacy and B43 in v16. Reading v16 addresses
    against a legacy file lands one row off and returns the wrong figure."""
    ctx = legacy_context
    assert ctx.sources["cost_grand_total"] == "Financial Worksheet!B42"
    assert ctx.sources["cost_after_discount"] == "Financial Worksheet!B43"
    assert ctx.num("cost_grand_total") == pytest.approx(315180.4257, abs=CENTS)
    assert ctx.num("cost_after_discount") == pytest.approx(272023.0287, abs=CENTS)


def test_legacy_reads_its_overrides_from_the_chassis_namespace(legacy_context):
    """The top-level `legacy:` key is gone; every address it used to supply now
    comes from `chassis.legacy_food4less` and must land in exactly the same
    place. This is the provenance gate on the override refactor - the figures
    are pinned elsewhere, but a silently-empty override block would read the
    v16 base map and still produce plausible numbers."""
    ctx = legacy_context
    assert ctx.sources["cost_grand_total"] == "Financial Worksheet!B42"
    assert ctx.sources["evolv_ports"] == "Internal Summary!G4"
    assert ctx.sources["svc_contract_years"] == "Internal Summary!G5"
    assert ctx.sources["cc_l2_credit_month"] == "Financial Worksheet!K46"
    assert ctx.tables["consolidated_cashflow"][0]["_row"] == 4


def test_legacy_itc_is_present_and_included_in_net(legacy_context):
    """Legacy B7 is =B3+(B4+B5+B6): a different formula, so ITC is inside."""
    ctx = legacy_context
    assert ctx.num("roi_itc") == pytest.approx(81606.9086, abs=CENTS)
    assert ctx.raw["has_itc"] is True
    expected = (ctx.num("roi_total_costs_upfront") + ctx.num("roi_carbon_credits")
                + ctx.num("roi_itc") + ctx.num("roi_evse_revenues"))
    assert ctx.num("roi_net_revenues") == pytest.approx(expected, abs=CENTS)


def test_legacy_error_field_is_not_emitted(legacy_context):
    """Legacy FW!B19 is a live #ERROR! (='Internal Summary'!#REF!). It maps to
    an emit:false token, so the "no token starts with #" gate must skip it
    rather than abort the acceptance run."""
    ctx = legacy_context
    assert "cost_service_list_price" not in ctx.emitted
    emitted_strings = [
        v for k, v in ctx.raw.items()
        if k in ctx.emitted and isinstance(v, str)
    ]
    assert not any(s.strip().startswith("#") for s in emitted_strings)


def test_legacy_cashflow_header_is_one_row_higher(legacy_context):
    rows = legacy_context.tables["consolidated_cashflow"]
    assert len(rows) == 11                     # year 0 through 10
    assert rows[0]["_row"] == 4                # legacy data starts at row 4
    assert rows[0]["cumulative"] == pytest.approx(-272023.0287, abs=1.0)


def test_legacy_carbon_divisor_is_reported(legacy_context):
    """The legacy workbook HAS the divisor defect, and now says so.

    This used to assert the opposite, on the reasoning that legacy is fixed
    10-year and `/(5*12)` is "what the reference proposal prints". That reasoning
    exempted the file from a check it fails. `Financial Worksheet!B4` holds ten
    years of carbon and `Cashflow!G3` spreads it over five, so the 60-month
    financing tables carry 2x the real monthly carbon - $3,640.99 where the true
    rate is $1,820.50.

    We still reproduce the reference document exactly, because we read cached
    values and the reference has the defect baked in. What changed is that the QA
    report now names it instead of staying silent. Fix `G3` to
    `='Financial Worksheet'!$B$4/(10*12)` in the workbook and this goes quiet.
    """
    codes = {f.code for f in legacy_context.findings}
    assert "cashflow-carbon-doubled" in codes


# --------------------------------------------------------------------------
# rejection
# --------------------------------------------------------------------------


def test_unsupported_workbook_is_rejected_cleanly(field_map, tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "not-a-calculator.xlsx"
    wb = Workbook()
    wb.active["A1"] = "hello"
    wb.save(path)

    with pytest.raises((FingerprintMismatch, UnknownEngine)) as exc:
        engine_mod.load(path, field_map)
    assert "supported" in exc.value.message.lower()
    assert "Traceback" not in exc.value.message
