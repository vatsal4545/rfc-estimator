"""Phase 2: proof that the Section 3B recompute is faithful.

The whole point of `baseline.py` is that it reproduces the specimen's Section 3B
*without touching the colliding block*. If these numbers match, the recompute is
correct and both engines will print the same Section 3B.
"""

from __future__ import annotations

import openpyxl
import pytest

from ev_proposal_agent import engine as engine_mod
from ev_proposal_agent.baseline import (
    detect_outage_runs,
    parse_window,
    recompute_level,
)
from ev_proposal_agent.extract import build_context
from ev_proposal_agent.findings import Findings

from .conftest import FOOD4LESS_OPERATOR_INPUTS


# --------------------------------------------------------------------------
# the acceptance figures
# --------------------------------------------------------------------------


def test_reproduces_food4less_section_3b(legacy_context):
    """Every printed figure in the reference proposal's Section 3B."""
    ctx = legacy_context
    l2 = ctx.raw["_baseline_l2"]
    l3 = ctx.raw["_baseline_l3"]

    # occupancy 15.3% / 43.0%
    assert l2.stall_occupancy == pytest.approx(0.1532, abs=1e-4)
    assert l3.stall_occupancy == pytest.approx(0.4304, abs=1e-4)

    # charger rating 7.06 / 49.00
    assert l2.charger_rating_kw == pytest.approx(7.056, abs=1e-3)
    assert l3.charger_rating_kw == pytest.approx(49.00, abs=1e-3)

    # profit per year $269.50 / $12,848.62
    assert l2.profit_per_year == pytest.approx(269.4976901, abs=0.01)
    assert l3.profit_per_year == pytest.approx(12848.61904, abs=0.01)

    # total yearly $13,118 and gross yearly $31,220
    assert ctx.num("hist_profit_yearly") == pytest.approx(13118.11673, abs=0.01)
    assert ctx.num("hist_gross_yearly") == pytest.approx(31219.68338, abs=0.01)


def test_intermediate_rows_match_the_specimen(legacy_context):
    """Not just the totals - every row of the chain."""
    l2 = legacy_context.raw["_baseline_l2"]
    l3 = legacy_context.raw["_baseline_l3"]

    assert l2.stalls_used_per_day == pytest.approx(0.766, abs=1e-3)
    assert l3.stalls_used_per_day == pytest.approx(2.152, abs=1e-3)
    assert l2.charging_hourly_pct == pytest.approx(0.2630833333, abs=1e-6)
    assert l3.charging_hourly_pct == pytest.approx(0.1343333333, abs=1e-6)
    assert l2.hours_per_stall_day == pytest.approx(3.157, abs=1e-3)
    assert l3.hours_per_stall_day == pytest.approx(1.612, abs=1e-3)
    assert l2.avg_delivered_pct == pytest.approx(0.3683553992, abs=1e-6)
    assert l3.avg_delivered_pct == pytest.approx(0.7025441375, abs=1e-6)
    assert l2.retail_rate == pytest.approx(0.5191032419, abs=1e-8)
    assert l3.retail_rate == pytest.approx(0.6988663043, abs=1e-8)
    assert l2.revenue_per_day == pytest.approx(3.262741784, abs=1e-6)
    assert l3.revenue_per_day == pytest.approx(83.45860094, abs=1e-6)
    assert l2.revenue_per_30day == pytest.approx(97.88225352, abs=1e-5)
    assert l3.revenue_per_30day == pytest.approx(2503.758028, abs=1e-4)
    assert l2.net_profit_month == pytest.approx(22.45814085, abs=1e-6)
    assert l3.net_profit_month == pytest.approx(1070.718254, abs=1e-5)
    assert l2.out_of_service_share == pytest.approx(0.58, abs=1e-9)
    assert l3.out_of_service_share == pytest.approx(0.28, abs=1e-9)


def test_recompute_matches_the_colliding_block_it_replaces(legacy_path, legacy_context):
    """The recompute must equal `Updated!B5:C27` in the one workbook where that
    block is safe to read. That is what proves reading it was never necessary."""
    wb = openpyxl.load_workbook(legacy_path, data_only=True)
    ws = wb["Updated Chargers Revenue Calcul"]
    l2 = legacy_context.raw["_baseline_l2"]
    l3 = legacy_context.raw["_baseline_l3"]

    for attr, row in [
        ("stall_occupancy", 8), ("stalls_used_per_day", 9), ("charging_hourly_pct", 10),
        ("hours_per_stall_day", 11), ("charger_rating_kw", 12), ("avg_delivered_pct", 13),
        ("retail_rate", 14), ("revenue_per_day", 15), ("revenue_per_30day", 16),
        ("kwh_per_day", 17), ("net_profit_month", 19), ("profit_per_year", 20),
    ]:
        assert getattr(l2, attr) == pytest.approx(float(ws[f"B{row}"].value), rel=1e-9), \
            f"L2 {attr} vs Updated!B{row}"
        assert getattr(l3, attr) == pytest.approx(float(ws[f"C{row}"].value), rel=1e-9), \
            f"L3 {attr} vs Updated!C{row}"

    assert legacy_context.num("hist_profit_yearly") == pytest.approx(
        float(ws["B24"].value), rel=1e-9)
    assert legacy_context.num("hist_gross_yearly") == pytest.approx(
        float(ws["B27"].value), rel=1e-9)
    wb.close()


def test_reconciliation_passes_on_the_specimen(legacy_context):
    """Updated!B22 vs C22: modelled revenue over the window vs actual.
    Both are $92,358.23."""
    codes = {f.code for f in legacy_context.findings}
    assert "baseline-reconciliation" not in codes


def test_profit_is_invariant_to_the_port_count(field_map, legacy_path):
    """Ports cancel out of the chain, so they cannot change revenue or profit.

        stalls_used_per_day = ports * (J12 / ports)      = J12
        rating * delivered_pct = rating * (J14 / rating) = J14

    This is worth pinning down because it is counter-intuitive and it is why
    the reconciliation cannot be used as a port-count check.
    """
    bad = dict(FOOD4LESS_OPERATOR_INPUTS, existing_ports_l3=12)
    with engine_mod.load(legacy_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=bad)

    assert ctx.num("hist_profit_yearly") == pytest.approx(13118.11673, abs=0.01)
    assert ctx.num("hist_gross_yearly") == pytest.approx(31219.68338, abs=0.01)
    # ...but the displayed occupancy does move, which is the real risk.
    assert ctx.raw["_baseline_l3"].stall_occupancy == pytest.approx(2.152 / 12, abs=1e-6)


def test_profit_is_invariant_to_nameplate_kw(field_map, legacy_path):
    """Same cancellation: `rating` multiplies in and divides out again."""
    bad = dict(FOOD4LESS_OPERATOR_INPUTS, existing_nameplate_l3_kw=350)
    with engine_mod.load(legacy_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=bad)
    assert ctx.num("hist_profit_yearly") == pytest.approx(13118.11673, abs=0.01)


def test_wrong_port_count_is_caught_against_the_station_note(field_map, legacy_path):
    """`Historical Data!I15` records the real station count. Since the maths
    cannot catch a bad port count, this note is what does."""
    bad = dict(FOOD4LESS_OPERATOR_INPUTS, existing_ports_l3=12)
    with engine_mod.load(legacy_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs=bad)
    codes = {f.code for f in ctx.findings}
    assert "port-count-disagrees-with-sheet" in codes


def test_specimen_l2_port_count_disagrees_with_its_own_note(legacy_context):
    """A real discrepancy in the source data, surfaced rather than hidden: the
    note records 6 Level 2 station IDs, the calculator was driven with 5 ports.
    Five ports across six stations is not a whole number either way, so it
    cannot be a dual-port arrangement. That choice is what produces the 15.3%
    occupancy the reference prints instead of 12.8%."""
    findings = [f for f in legacy_context.findings
                if f.code == "port-count-disagrees-with-sheet"]
    assert len(findings) == 1
    assert "L2: 5 ports entered against 6 station IDs" in findings[0].message


# --------------------------------------------------------------------------
# the window
# --------------------------------------------------------------------------


def test_window_is_parsed_from_the_formula(field_map, legacy_path):
    """`Historical Data!J4` is `=ROWS(A8:A42)`. Rows 8-42, not 9-43.

    Guessing from the data picks up the trailing partial 2026-07 (which still
    carries $1,040) and shifts the entire window by one month.
    """
    with engine_mod.load(legacy_path, field_map) as lw:
        start, end = parse_window(lw, field_map["full_history"]["monthly_series"],
                                  Findings())
    assert (start, end) == (8, 42)


def test_window_totals_exclude_the_stubs_and_the_partial(legacy_context):
    """`Historical Data!D44` sums all 40 rows and reads $93,453.35. The window
    total used everywhere downstream is $92,358.23."""
    assert legacy_context.num("hist_total_revenue") == pytest.approx(92358.23, abs=0.5)
    assert legacy_context.num("hist_total_revenue") != pytest.approx(93453.35, abs=0.5)


def test_window_covers_thirty_five_months(legacy_context):
    points = legacy_context.tables["monthly_series"]
    assert len(points) == 35
    assert points[0]["label"] == "Aug 2023"
    assert points[-1]["label"] == "Jun 2026"
    assert legacy_context.raw["history_window_label"] == "Aug 2023 to Jun 2026"


def test_dcfc_revenue_share(legacy_context):
    assert legacy_context.num("dcfc_revenue_share") == pytest.approx(0.962, abs=0.002)


# --------------------------------------------------------------------------
# outage detection
# --------------------------------------------------------------------------


def test_outage_is_detected_not_hardcoded(legacy_context):
    """Jan and Feb 2026 collapsed to $38.80 and $121.18 of Level 3 revenue.
    The rule is "more than 80% below the trailing three-month mean", so it will
    also catch an outage on a site whose dates nobody knows in advance."""
    points = legacy_context.tables["monthly_series"]
    runs = detect_outage_runs(points)
    assert runs, "the Jan-Feb 2026 outage should be found"
    start, end = runs[0]
    assert points[start]["label"] == "Jan 2026"
    assert points[end]["label"] == "Feb 2026"
    assert legacy_context.raw["outage_window_label"] == "Jan 2026 and Feb 2026"


def test_no_outage_flagged_on_a_steady_series():
    points = [{"label": f"M{i}", "l3_rev": 1000.0 + i} for i in range(12)]
    assert detect_outage_runs(points) == []


# --------------------------------------------------------------------------
# guards
# --------------------------------------------------------------------------


def test_zero_ports_yields_a_zero_baseline_not_a_crash():
    """A site can genuinely have had no chargers of one level - Showcase Liquor
    ran Level 2 only. That is a real baseline of nothing, so it must produce
    zeroes rather than raise or divide by zero."""
    result = recompute_level(
        level="L3", ports=0, nameplate_kw=50, derate=0.98,
        avg_stalls_used_day=0, avg_charging_hours=0,
        avg_delivered_kw=0, effective_rate=0, unproductive_share=0,
    )
    assert result.profit_per_year == 0.0
    assert result.revenue_per_day == 0.0
    assert result.stall_occupancy == 0.0
    # the rating is still meaningful - it is nameplate x de-rate, not a division
    assert result.charger_rating_kw == pytest.approx(49.0)


def test_ports_prefill_from_the_workbooks_own_historical_column(field_map, legacy_path):
    """With no operator input at all, Section 3B still builds.

    `Updated!B6`/`C6` hold the EXISTING stall counts under the legacy engine, so
    there is no reason to make anyone retype them. Note this is NOT
    `Internal Summary!G3`, which is the count of chargers being installed and
    would put the occupancy several times too high.
    """
    with engine_mod.load(legacy_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs={})

    assert ctx.raw["existing_ports_l2"] == 5
    assert ctx.raw["existing_ports_l3"] == 5
    assert ctx.raw["existing_nameplate_l2_kw"] == pytest.approx(7.2)
    assert ctx.raw["has_baseline"] is True
    assert ctx.num("hist_profit_yearly") == pytest.approx(13118.11673, abs=0.01)


def test_ports_are_not_prefilled_under_the_scenario_grid(field_map, v16_paths):
    """`Updated!B6` means something else entirely there - the Standard-Low NEW
    stall quantity. Prefilling from it is exactly the collision this program
    exists to avoid."""
    with engine_mod.load(v16_paths, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs={})
    assert ctx.raw.get("existing_ports_l2") is None


def test_section_3b_is_skipped_when_ports_are_genuinely_zero(field_map, legacy_path):
    """An explicit zero from the operator wins over the workbook prefill, and
    with no ports at all there is no Section 3B."""
    with engine_mod.load(legacy_path, field_map) as lw:
        ctx = build_context(lw, field_map, operator_inputs={
            "existing_ports_l2": 0, "existing_ports_l3": 0})
    assert ctx.raw["has_baseline"] is False
    assert "baseline-ports-missing" in {f.code for f in ctx.findings}


def test_baseline_table_has_every_reference_row(legacy_context):
    rows = legacy_context.tables["historical_baseline"]
    labels = [r["label"] for r in rows]
    assert len(rows) == 17
    assert labels[0] == "Chargers"
    assert "Total Profit per Year" in labels
    assert all(r["l2"] and r["l3"] for r in rows)
