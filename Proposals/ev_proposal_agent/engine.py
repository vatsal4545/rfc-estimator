"""Which revenue engine is wired, which scenario, and over what horizon.

This runs before anything else reads a number. Two incompatible revenue models
share the sheet name `Updated Chargers Revenue Calcul` **and the same cell
addresses**, so reading the wrong one returns plausible wrong numbers with no
error at all:

    Updated!B19   legacy: L2 net profit per MONTH
                  v16:    L2 total profit per YEAR
    Updated!G12   legacy: L2 idle-fee profit per year
                  v16:    Standard-Low L3 240 kW charger rating

Probe `Updated!B3` and branch. Never assume.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from .errors import (
    FingerprintMismatch,
    UnknownEngine,
    WorkbookNotCalculated,
    WorkbookNotReadable,
)
from .findings import Findings
from .resolve import is_missing, sheet_has_cached_values

SCENARIO_GRID = "scenario_grid"
BASELINE_VS_PROJECTED = "baseline_vs_projected"

UCRC = "Updated Chargers Revenue Calcul"
FW = "Financial Worksheet"
HISTORICAL_DATA = "Historical Data"
TEN_YEAR_SHEET = "10 Year Projection Comparison"
DLL = "DLL Schedule"
CASHFLOW = "Cashflow"

# 'Financial Worksheet'!I6 points at the totals cell of the wired scenario.
_SCENARIO_BY_TOTALS_CELL = {
    "J24": "Standard-Low",
    "P24": "Standard-Medium",
    "Q24": "Standard-High",
}

_TOTALS_CELL_RE = re.compile(
    r"Updated\s+Chargers\s+Revenue\s+Calcul'?!\$?([A-Z]{1,2})\$?(\d+)", re.IGNORECASE
)


@dataclass
class Detection:
    """Everything downstream needs to know before reading a single figure.

    `chassis` and `engine` are INDEPENDENT. The chassis says where the rows are
    on Financial Worksheet / Internal Summary / INPUT SHEET / DLL / Cashflow.
    The engine says which revenue model `Updated Chargers Revenue Calcul` holds.
    The production workbook is a v16 chassis carrying the legacy engine, so
    anything that assumes the two agree rejects the real input.
    """

    chassis: str
    engine: str
    scenario: str | None
    projection_years: int
    has_history: bool
    has_financing: bool
    ten_year_sheet_present: bool
    ten_year_sheet_readable: bool = False
    findings: Findings = field(default_factory=Findings)
    # `field_map['chassis'][chassis]` - the per-section overrides that replace
    # the base (v16) map for this layout. Empty for v16, which IS the base map.
    chassis_spec: dict = field(default_factory=dict)

    @property
    def is_legacy_chassis(self) -> bool:
        """Diagnostic only. Nothing keys behaviour off this any more.

        Row layout is decided by `chassis_spec`, because the third chassis
        (`v15_no_itc`) takes the legacy cost rows and the v16 EVOLV rows at the
        same time. A boolean cannot express that, and asking it produced an
        EVOLV block read one row low.
        """
        return self.chassis == "legacy_food4less"

    def chassis_opt(self, section: str, key: str, default=None):
        """One option out of this chassis's override block for one map section."""
        return (self.chassis_spec.get(section) or {}).get(key, default)

    @property
    def is_legacy_engine(self) -> bool:
        """Revenue model: `Updated` holds baseline-vs-projected with idle fees,
        not the seven-column scenario grid."""
        return self.engine == BASELINE_VS_PROJECTED

    def summary_pairs(self) -> list[tuple[str, str]]:
        """(label, value) for the app. Structured, so no caller has to re-split
        a padded string - "Projection horizon 5 years" has a single space and
        splitting on two put the whole line in the label column."""
        engine_label = {
            SCENARIO_GRID: "scenario grid (Low / Med / High tiers)",
            BASELINE_VS_PROJECTED: "historical vs projected (with idle fees)",
        }.get(self.engine, self.engine)
        ten_year = (
            "present and readable" if self.ten_year_sheet_readable
            else "present but MIS-WIRED, not read" if self.ten_year_sheet_present
            else "absent"
        )
        return [
            ("Workbook layout", self.chassis),
            ("Revenue model", engine_label),
            ("Scenario", self.scenario or "n/a for this revenue model"),
            ("Projection horizon", f"{self.projection_years} years"),
            ("Financing", "active" if self.has_financing else "none"),
            ("Historical data", "present" if self.has_history else "absent"),
            ("10-year sheet", ten_year),
        ]

    def summary_lines(self) -> list[str]:
        """What the CLI and the QA report show."""
        engine_label = {
            SCENARIO_GRID: "scenario grid (Low / Med / High tiers)",
            BASELINE_VS_PROJECTED: "historical vs projected (with idle fees)",
        }.get(self.engine, self.engine)
        return [
            f"Workbook layout    {self.chassis}",
            f"Revenue model      {engine_label}",
            f"Scenario           {self.scenario or 'n/a for this revenue model'}",
            f"Projection horizon {self.projection_years} years",
            f"Financing          {'active' if self.has_financing else 'none'}",
            f"Historical data    {'present' if self.has_history else 'absent'}",
            f"10-year sheet      " + (
                "present and readable" if self.ten_year_sheet_readable
                else "present but MIS-WIRED - not read" if self.ten_year_sheet_present
                else "absent"
            ),
        ]


@dataclass
class LoadedWorkbook:
    """Both loads of the same file, plus the detection result.

    Scenario detection reads a *formula*, and cached values are what everything
    else reads, so both loads are mandatory. Keeping them together stops anyone
    from accidentally reading a formula string where a number belonged.
    """

    path: Path
    values: Workbook      # data_only=True  - cached results
    formulas: Workbook    # data_only=False - formula text
    detection: Detection

    def close(self) -> None:
        self.values.close()
        self.formulas.close()

    def __enter__(self) -> "LoadedWorkbook":
        return self

    def __exit__(self, *exc) -> None:
        self.close()


# --------------------------------------------------------------------------


def _probe_failure(wb: Workbook, probe: dict) -> str | None:
    """None if the probe passes, else a human-readable reason it didn't."""
    sheet, cell, expected = probe["sheet"], probe["cell"], probe["equals"]
    if sheet not in wb.sheetnames:
        return f"no sheet named '{sheet}'"
    got = wb[sheet][cell].value
    if str(got).strip() != expected:
        return f"{sheet}!{cell} should read '{expected}' but reads '{got}'"
    return None


def identify_chassis(wb: Workbook, meta: dict) -> tuple[str, dict]:
    """Return the first chassis whose probes all pass.

    This identifies the ROW LAYOUT only. It says nothing about which revenue
    model the Updated sheet holds - that is `detect_engine`, and the two vary
    independently. Keeping this table-driven is what lets a future v17 be added
    as data rather than as branches scattered through the extractor.
    """
    variants: dict = meta.get("fingerprints") or {}
    if not variants:  # older map shape, single unnamed fingerprint
        for probe in meta.get("fingerprint", []):
            reason = _probe_failure(wb, probe)
            if reason:
                raise FingerprintMismatch(
                    "This does not look like a supported calculator workbook.",
                    detail=reason,
                )
        return "v16", {"required_sheets": meta.get("required_sheets", [])}

    reasons: list[str] = []
    for name, spec in variants.items():
        failures = [r for p in spec.get("probes", []) if (r := _probe_failure(wb, p))]
        if not failures:
            return name, spec
        reasons.append(f"{name}: {failures[0]}")

    raise FingerprintMismatch(
        "This does not look like a supported calculator workbook.",
        detail="; ".join(reasons),
    )


def check_required_sheets(wb: Workbook, required: list[str]) -> None:
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        raise FingerprintMismatch(
            "The workbook is missing sheets this program needs: "
            + ", ".join(missing) + ".",
            detail="Is this a v16 workbook?",
        )


def detect_engine(wb: Workbook, spec: dict) -> str:
    """Read `Updated!B3` and map it. Startswith, because the legacy banner
    carries a date range: `HISTORICAL ACTUAL (Aug 2023 - Jun 2026)`."""
    probe = spec["probe"]
    sheet = probe["sheet"]
    if sheet not in wb.sheetnames:
        raise UnknownEngine(
            f"The workbook has no '{sheet}' tab, so the revenue model cannot "
            "be identified."
        )
    raw = wb[sheet][probe["cell"]].value
    if is_missing(raw):
        raise UnknownEngine(
            f"{sheet}!{probe['cell']} is empty, so the revenue model cannot be "
            "identified. Open the workbook in Excel, press Ctrl+Alt+F9, save "
            "and try again."
        )
    text = str(raw).strip()

    for key, engine in spec["values"].items():
        if text == key or text.startswith(key):
            return engine

    raise UnknownEngine(
        "Unrecognised revenue model on the Updated Chargers Revenue Calcul "
        f'tab. Cell {probe["cell"]} holds "{text}". This file is not a '
        "supported calculator version."
    )


def detect_scenario(formulas: Workbook, findings: Findings, *, cell: str = "I6") -> str:
    """Parse the FORMULA of the scenario pointer to find the wired scenario.

    Not the value - the value is just a number that could belong to any column.
    v16 ships `='Updated Chargers Revenue Calcul'!J24` (Standard-Low) at `I6`.

    The cell MOVES with the chassis. On v15_no_itc the charger-cashflow block
    sits a row higher, so the pointer is at `I5` and `I6` holds `=I5*1.125` -
    no reference to the Updated tab, and a silent fall back to Standard-Low on a
    workbook that might be wired to Medium or High. Chassis declares it.

    No scenario picker is exposed anywhere in the app: the cashflow tables are
    wired to exactly one scenario, so letting the operator choose a different
    one would make Section 5 contradict Section 7A.
    """
    formula = formulas[FW][cell].value
    if not isinstance(formula, str) or "!" not in formula:
        findings.warn(
            "scenario-unreadable",
            "Could not read the scenario link; assuming Standard-Low. "
            f"{FW}!{cell} holds {formula!r} instead of a reference to the "
            "Updated Chargers Revenue Calcul tab.",
            where=f"{FW}!{cell}",
        )
        return "Standard-Low"

    match = _TOTALS_CELL_RE.search(formula)
    if not match:
        findings.warn(
            "scenario-unparsed",
            f"Could not parse the scenario link {formula!r}; assuming "
            "Standard-Low.",
            where=f"{FW}!{cell}",
        )
        return "Standard-Low"

    totals_cell = f"{match.group(1).upper()}{match.group(2)}"
    scenario = _SCENARIO_BY_TOTALS_CELL.get(totals_cell)
    if scenario is None:
        findings.warn(
            "scenario-unknown-cell",
            f"{FW}!{cell} points at {totals_cell}, which is not one of the three "
            "scenario totals cells (J24, P24, Q24). Assuming Standard-Low.",
            where=f"{FW}!{cell}",
        )
        return "Standard-Low"

    if scenario != "Standard-Low":
        findings.warn(
            "scenario-rewired",
            f"The model is wired to {scenario}, not the shipped Standard-Low. "
            "Every figure in the proposal will use that scenario.",
            where=f"{FW}!{cell}",
        )
    return scenario


def detect_horizon(values: Workbook, findings: Findings, *, fixed: int | None = None) -> int:
    """`Financial Worksheet!B10` = `=IF(B9="10 Year",10,5)`.

    A chassis with no B9/B10 toggle declares `horizon.fixed` in the map instead
    of being recognised by name.
    """
    if fixed:
        return int(fixed)
    raw = values[FW]["B10"].value
    if isinstance(raw, (int, float)) and int(raw) in (5, 10):
        return int(raw)
    toggle = values[FW]["B9"].value
    findings.warn(
        "horizon-unreadable",
        f"{FW}!B10 holds {raw!r} (the B9 toggle reads {toggle!r}); defaulting "
        "to a 5-year projection.",
        where=f"{FW}!B10",
    )
    return 5


def _detect_financing(values: Workbook, findings: Findings) -> bool:
    """A zero or blank `Loan_Amount` means no financing; the DLL summary cells
    return `""` in that state and Sections 17 / 17A-D are suppressed."""
    if DLL not in values.sheetnames:
        return False
    amount = values[DLL]["D5"].value
    if is_missing(amount) or not isinstance(amount, (int, float)):
        return False
    return float(amount) > 0


# --------------------------------------------------------------------------


def _open_both(path: Path) -> tuple[Workbook, Workbook]:
    """Both loads of the workbook, with a recovery pass and a real diagnosis.

    A dropped file that cannot be read used to surface as a bare
    "permission denied", which tells the operator nothing about which of the
    three usual causes it is (open in Excel / OneDrive online-only / dragged out
    of Outlook).

    The recovery is worth trying first: Windows will normally let us COPY a file
    that Excel holds open, even when opening it in place is refused, so copying
    to temp and reading the copy fixes the commonest case outright rather than
    telling the operator to go and close Excel. The copy is read-only to us and
    the original is never touched.
    """
    try:
        return (load_workbook(path, data_only=True, read_only=False),
                load_workbook(path, data_only=False, read_only=False))
    except (PermissionError, OSError) as first:
        import shutil
        import tempfile

        try:
            scratch = Path(tempfile.mkdtemp(prefix="evproposal-")) / path.name
            shutil.copy2(path, scratch)
            return (load_workbook(scratch, data_only=True, read_only=False),
                    load_workbook(scratch, data_only=False, read_only=False))
        except (PermissionError, OSError, shutil.Error):
            raise WorkbookNotReadable.diagnose(path, first) from first


def stale_cached_values(values: Workbook) -> list[str]:
    """Discrepancies proving the cached results predate the inputs.

    openpyxl reads only what the spreadsheet cached when it was last saved. A
    workbook that was calculated once and then edited by a writer that is not
    Excel keeps its OLD answers: the equipment and cost inputs are typed in and
    plainly visible, and the formulas that consume them still hold the values
    from before those inputs existed.

    That reads as a perfectly valid workbook for a project that costs nothing,
    so the app produced a $0 proposal and suppressed the financing sections
    with no complaint. Nothing looked wrong - the file is full of numbers.

    Detected semantically, by checking identities the workbook must satisfy,
    NOT by inspecting the package. `fullCalcOnLoad="1"` looked like a clean
    signal and is not: openpyxl sets it on every save, so any programmatically
    written workbook carries it - including this project's own synthetic test
    fixture, which is correctly valued and must load.

    Returns a list of human-readable discrepancies; empty means consistent.
    """
    problems: list[str] = []

    # 1. Equipment priced at nothing. `line_total` is a formula over a
    #    price-book lookup; a SKU with a quantity and no line total means the
    #    lookup has never run against this row.
    ins = values["INPUT SHEET"] if "INPUT SHEET" in values.sheetnames else None
    if ins is not None:
        for row in range(8, 38):
            sku = ins[f"E{row}"].value
            qty = ins[f"I{row}"].value
            if not sku or not isinstance(qty, (int, float)) or qty <= 0:
                continue
            total = ins[f"K{row}"].value
            if total is None or total == "":
                problems.append(
                    f"INPUT SHEET!E{row} is {sku!r} x {qty:g} but K{row} "
                    f"(line total) has never been computed")
            elif total == 0:
                problems.append(
                    f"INPUT SHEET!E{row} is {sku!r} x {qty:g} but K{row} "
                    f"(line total) is 0")

    # 2. Customer price out of step with list price and discount. Every cost
    #    row on Internal Summary is D = B - (B x C). If B is filled and D
    #    disagrees, D is an answer to an older question.
    if "Internal Summary" in values.sheetnames:
        ws = values["Internal Summary"]
        for row in range(3, 29):
            listed = ws[f"B{row}"].value
            discount = ws[f"C{row}"].value or 0
            customer = ws[f"D{row}"].value
            if not isinstance(listed, (int, float)) or listed == 0:
                continue
            if not isinstance(discount, (int, float)):
                continue
            if not isinstance(customer, (int, float)):
                continue
            expected = listed - (listed * discount)
            if abs(customer - expected) > 0.01:
                label = ws[f"A{row}"].value or f"row {row}"
                problems.append(
                    f"Internal Summary!D{row} ({label}) is "
                    f"{customer:,.2f} but B{row} - (B{row} x C{row}) is "
                    f"{expected:,.2f}")

    return problems


def load(path: str | Path, field_map: dict) -> LoadedWorkbook:
    """Open a workbook, verify it, and detect engine / scenario / horizon.

    Everything that reads a figure goes through the object this returns.
    """
    path = Path(path)
    findings = Findings()

    values, formulas = _open_both(path)

    meta = field_map["meta"]
    chassis, fingerprint_spec = identify_chassis(values, meta)
    check_required_sheets(
        values, fingerprint_spec.get("required_sheets") or meta["required_sheets"]
    )

    # The row/address overrides for this layout. v16 has no entry: it IS the
    # base map, and an entry for it would mean the base map is wrong.
    chassis_spec = (field_map.get("chassis") or {}).get(chassis) or {}

    if not sheet_has_cached_values(values[FW]):
        raise WorkbookNotCalculated(detail=f"no cached numbers on '{FW}'")

    stale = stale_cached_values(values)
    if stale:
        shown = "; ".join(stale[:3])
        more = f" (+{len(stale) - 3} more)" if len(stale) > 3 else ""
        raise WorkbookNotCalculated(
            "This workbook has not been recalculated since it was last edited. "
            "Its inputs are filled in, but the formulas that price them still "
            "hold their previous answers, so every cost, total and loan figure "
            "would be wrong. Open it in Excel, press Ctrl+Alt+F9, save, and "
            "try again.",
            detail=shown + more,
        )

    # Independent of the chassis. The production workbook is a v16 chassis
    # carrying the legacy engine, so these are read separately and never
    # cross-checked against each other.
    engine = detect_engine(values, field_map["engine"])
    legacy_engine = engine == BASELINE_VS_PROJECTED

    # `horizon.fixed` means this chassis has no B9/B10 toggle at all.
    fixed_horizon = (chassis_spec.get("horizon") or {}).get("fixed")

    scenario_cell = (chassis_spec.get("scenario_pointer") or {}).get("cell", "I6")
    scenario = (None if legacy_engine
                else detect_scenario(formulas, findings, cell=scenario_cell))
    projection_years = detect_horizon(values, findings, fixed=fixed_horizon)

    has_history = HISTORICAL_DATA in values.sheetnames
    if not has_history:
        findings.info(
            "history-absent",
            f"No '{HISTORICAL_DATA}' sheet. Sections 3, 3A and 3B, the "
            "Section 5 baseline comparison and the historical revenue chart "
            "will be suppressed and the document renumbered. Paste the sheet "
            "in from the utilization workbook to include them.",
        )

    # The 10-year sheet's formulas are bound to the legacy Updated addresses.
    # That makes it CORRECT when the legacy engine is present (the normal
    # production case) and silently WRONG under the scenario grid, where the
    # same addresses hold charger ratings and annual totals instead.
    ten_year_present = TEN_YEAR_SHEET in values.sheetnames
    ten_year_readable = ten_year_present and legacy_engine
    if ten_year_present and not legacy_engine:
        findings.error(
            "ten-year-sheet-miswired",
            f"'{TEN_YEAR_SHEET}' is present but its formulas are bound to the "
            "historical-vs-projected revenue model, while this workbook's "
            "Updated Chargers Revenue Calcul tab holds the scenario grid. At "
            "those same addresses the grid keeps charger ratings where idle "
            "fees belong and annual figures where monthly ones do, so the "
            "sheet fills with plausible wrong numbers and reports no error. It "
            "has not been read; the ten-year figures are derived from the "
            "Financial Worksheet instead. Delete the sheet to clear this.",
            where=TEN_YEAR_SHEET,
        )

    has_financing = _detect_financing(values, findings)

    detection = Detection(
        chassis=chassis,
        engine=engine,
        scenario=scenario,
        projection_years=projection_years,
        has_history=has_history,
        has_financing=has_financing,
        ten_year_sheet_present=ten_year_present,
        ten_year_sheet_readable=ten_year_readable,
        findings=findings,
        chassis_spec=chassis_spec,
    )

    _check_carbon_divisor(formulas, findings, projection_years, has_financing)

    return LoadedWorkbook(path=path, values=values, formulas=formulas, detection=detection)


# `Cashflow!G3` should be ='Financial Worksheet'!$B$4/(<horizon> * 12). Capture
# whatever the divisor actually is so it can be judged rather than assumed.
_CARBON_DIVISOR_RE = re.compile(r"/\s*\(\s*([^)]*?)\s*\*\s*12\s*\)")


def _check_carbon_divisor(
    formulas: Workbook, findings: Findings, projection_years: int, has_financing: bool
) -> None:
    """Is `Cashflow!G3`'s divisor keyed to the horizon, or hardcoded?

    `B4` is the carbon total over the WHOLE horizon - it is itself
    `B10 x` an annual figure - so the monthly rate is `B4 / (B10 * 12)` and is
    horizon-invariant. The loan is 60 months whatever the view period, so these
    tables should not move at all when the toggle changes.

    The shipped calculator writes `/(5*12)`, which is right only at a 5-year
    horizon. At 10 years it spreads ten years of carbon over five and doubles
    the monthly figure.

    This used to be keyed to `fixed_horizon is None`, i.e. "does this chassis
    have a B9/B10 toggle", on the reasoning that a fixed-10-year workbook had
    always printed `/60` and nothing had regressed. That exempted the two
    toggleless chassis from a check they fail: Food4Less and both Marriott files
    carry the identical 2x error and generated without complaint. So the test is
    now the formula itself, on every chassis - and a workbook whose divisor has
    been corrected to reference `B10` passes at either setting, which is what
    lets an operator fix the sheet and move on.

    Read from the FORMULAS workbook: openpyxl hands back cached values from the
    other load, and a cached number cannot tell you what divided it.
    """
    if not has_financing or CASHFLOW not in formulas.sheetnames:
        return

    formula = formulas[CASHFLOW]["G3"].value
    if not isinstance(formula, str):
        return                      # a literal, or empty - nothing to judge

    # A divisor that references the horizon cell is correct by construction, at
    # any horizon, including ones nobody has added to B9 yet.
    if "B10" in formula.upper().replace("$", ""):
        return

    match = _CARBON_DIVISOR_RE.search(formula)
    if match is None:
        return                      # some other shape; not ours to second-guess
    try:
        divisor_years = float(match.group(1))
    except ValueError:
        return                      # a reference we do not recognise

    if abs(divisor_years - projection_years) < 0.001:
        return                      # hardcoded, but hardcoded correctly

    factor = projection_years / divisor_years if divisor_years else 0
    findings.error(
        "cashflow-carbon-doubled",
        f"{CASHFLOW}!G3:G62 divides the carbon total by ({match.group(1)}*12), "
        f"but this workbook is on a {projection_years}-year projection. "
        f"'Financial Worksheet'!B4 holds {projection_years} years of carbon, so "
        f"spreading it over {match.group(1)} years overstates the monthly figure "
        f"in the 60-month financing tables by {factor:g}x. Fix it at source - set "
        f"{CASHFLOW}!G3 to =\'Financial Worksheet\'!$B$4/(\'Financial "
        f"Worksheet\'!$B$10*12) and fill down to G62, which is correct at every "
        f"horizon and leaves a 5-year workbook unchanged to the cent.",
        where=f"{CASHFLOW}!G3",
    )
