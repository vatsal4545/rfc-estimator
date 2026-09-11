"""Workbook in, Word document out.

Order matters and is enforced here rather than left to the caller:

    detect -> extract -> plan sections -> validate -> charts -> render -> verify

Validation happens **before** anything is written. An ERROR-level finding means
no file appears at all, because a proposal with a wrong number in it is worse
than no proposal - it gets sent.
"""

from __future__ import annotations

import datetime as _dt
import json
import re
from pathlib import Path
from typing import Any

from docxtpl import DocxTemplate, InlineImage, RichText
from docx.shared import Inches

from . import engine as engine_mod
from .errors import OutputNotWritable, ProposalError, TemplateMismatch, ValidationFailed
from .extract import Context, build_context, load_field_map
from .findings import Findings
from .paths import PROPOSAL_TEMPLATE, output_dir
from .sections import build_plan

# Display width reserved in the template for each figure, in inches.
CHART_WIDTHS = {
    "chart_historical_revenue": 6.65,
    "chart_annual_operating_profit": 6.55,
    "chart_cumulative_cashflow": 5.85,
}

# Every `{%tr for %}` collection the template can reference.
LOOP_COLLECTIONS = (
    "toc_rows", "contents_rows", "cost_rows", "roi_rows", "equipment_rows",
    "infrastructure_rows", "consolidated_cashflow", "charger_cashflow",
    "full_history", "historical_baseline", "operating_model",
    "monthly_cashflow_1_25", "monthly_cashflow_26_50", "monthly_cashflow_51_60",
)

_UNSAFE_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def safe_filename(text: str, *, fallback: str = "Proposal") -> str:
    cleaned = _UNSAFE_FILENAME.sub("", str(text or "")).strip().strip(".")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned[:80] or fallback


def default_output_path(ctx: Context, *, directory: Path | None = None) -> Path:
    """`{site_name}_{job_number}_{date}.docx` in the operator's Documents.

    The job number is often absent - the INPUT SHEET client-info cell does not
    always carry one - so it is dropped from the name rather than left as an
    empty separator.
    """
    site = safe_filename(ctx.raw.get("site_name") or "Proposal")
    job = safe_filename(ctx.raw.get("job_number") or "")
    date = _dt.date.today().isoformat()
    stem = "_".join(part for part in (site, job, date) if part)
    return (directory or output_dir()) / f"{stem}.docx"


# --------------------------------------------------------------------------
# context assembly
# --------------------------------------------------------------------------


def build_render_context(ctx: Context, plan, tpl: DocxTemplate,
                         charts: dict[str, Path],
                         fm: dict | None = None) -> dict[str, Any]:
    """Flatten everything the template can reference into one dict.

    `fm` is the field map. Section 8's row set is declared there rather than in
    this module, so it is threaded in; it falls back to loading the map for the
    handful of tests that call this directly.
    """
    if fm is None:
        fm = load_field_map()
    from .narrative import build_all as build_narrative

    payload: dict[str, Any] = {}

    # Formatted strings for scalars; raw values stay available under `raw_`
    # for any template expression that needs to compare rather than print.
    payload.update(ctx.formatted)
    payload.update({f"raw_{k}": v for k, v in ctx.raw.items()
                    if not k.startswith("_")})

    # Booleans must be real booleans, not the strings `put()` formatted them to,
    # or `{% if has_itc %}` is true for "False".
    for flag in ("has_itc", "has_financing", "has_l2", "has_l3", "has_history",
                 "has_baseline"):
        payload[flag] = bool(ctx.raw.get(flag))

    payload.update(ctx.tables)
    payload.update(plan.flags())
    payload["sections"] = plan.sections
    payload["toc_rows"] = plan.toc_rows()
    payload["contents_rows"] = plan.contents_pairs()
    payload["section_number"] = plan.by_key
    payload.update(build_narrative(ctx))

    # Seed every loop collection, so one that is absent iterates zero times
    # rather than being filled with "" and iterating its characters.
    for name in LOOP_COLLECTIONS:
        payload.setdefault(name, [])

    payload["cost_rows"] = ctx.raw.get("_cost_rows") or []
    payload["infrastructure_rows"] = _infrastructure_rows(ctx, fm)
    payload["equipment_rows"] = _equipment_rows(ctx)
    payload["roi_rows"] = _roi_rows(ctx)

    for token, path in charts.items():
        width = CHART_WIDTHS.get(token, 6.5)
        payload[token] = InlineImage(tpl, str(path), width=Inches(width))
    for token in CHART_WIDTHS:
        payload.setdefault(token, "")

    _add_product_photos(payload, ctx, tpl)
    _add_cover_photo(payload, ctx, tpl)

    # One multiplier line per L3 tier being installed.
    payload["cc_multiplier_lines"] = _kpi_richtext(
        ctx.formatted.get("cc_multiplier_lines", "") or "n/a")
    return payload


# The Section 13 KPI figure style, copied from the run this replaces:
#   <w:rPr><w:b/><w:color w:val="08B3AD"/><w:sz w:val="32"/></w:rPr>
# The template tag uses `{{r ... }}`, which swaps the whole run, so the
# RichText has to carry the styling or the figure drops to plain body text.
KPI_FIGURE_STYLE = {"bold": True, "color": "08B3AD", "size": 32}


def _kpi_richtext(text: str) -> RichText:
    """A KPI figure, one line per value, in the tile's own style.

    Built run by run rather than handing RichText a string containing newlines.
    An unstyled `RichText("$4,300")` emits the text with no run wrapper, which
    lands loose inside the paragraph and Word discards it, so the tile renders
    empty. Styling every run also forces the wrapper to exist.
    """
    rich = RichText()
    for index, line in enumerate(text.split("\n")):
        if index:
            rich.add("\n", **KPI_FIGURE_STYLE)      # becomes <w:br/>
        rich.add(line, **KPI_FIGURE_STYLE)
    return rich


# The reference cover photo's display size. A replacement is sized to the same
# box so a portrait snapshot from a phone does not blow the cover layout apart.
COVER_PHOTO_SIZE_IN = (6.45, 4.84)


def _add_cover_photo(payload: dict, ctx: Context, tpl: DocxTemplate) -> None:
    """The operator's site photo for the cover, falling back to the stock shot.

    Constrained by width only. Fixing both dimensions would stretch anything
    that is not exactly 4:3, and a distorted building on page one is worse than
    a slightly short one.
    """
    from .paths import TEMPLATES_DIR

    chosen = ctx.raw.get("cover_photo_path")
    path = Path(chosen) if chosen else TEMPLATES_DIR / "assets" / "cover_default.jpeg"
    if not path.exists():
        payload["cover_photo"] = ""
        return
    payload["cover_photo"] = InlineImage(
        tpl, str(path), width=Inches(COVER_PHOTO_SIZE_IN[0])
    )


# Reference display sizes. Sized by HEIGHT so the tall thin Level 2 pedestal and
# the wide Level 3 cabinet sit on the same baseline when both are shown.
PRODUCT_PHOTOS = {
    "photo_level2": ("charger_level2.png", "has_l2", 2.48),
    "photo_level3": ("charger_level3.png", "has_l3", 2.56),
}


def _add_product_photos(payload: dict, ctx: Context, tpl: DocxTemplate) -> None:
    """Insert each charger photo only when that level is actually being installed.

    A Level-2-only site should not be shown a DC fast charger it is not buying.
    These used to be page-anchored drawings in the reference, which is why they
    landed on the body copy whenever a table changed length.
    """
    from .paths import TEMPLATES_DIR

    for token, (filename, flag, height_in) in PRODUCT_PHOTOS.items():
        path = TEMPLATES_DIR / "assets" / filename
        if payload.get(flag) and path.exists():
            payload[token] = InlineImage(tpl, str(path), height=Inches(height_in))
        else:
            payload[token] = ""


def _infrastructure_rows(ctx: Context, fm: dict) -> list[dict[str, Any]]:
    """Section 8 reprints part of the cost stack with an explanation column.

    The rows come from `infrastructure.named_rows` in the field map, in cost-stack
    order, plus an "Other / miscellaneous" row carrying the residual. This was a
    hardcoded six-token dict that dropped eight of the fourteen children of
    `cost_electrical_subtotal` - on Best Western it printed $84,486.52 of a
    $191,639.04 scope and showed "$0 or project-specific" for the remainder.

    Zero rows are dropped, so a site with no switchgear does not print a $0 line;
    whatever they would have contributed is zero and the total is unaffected.
    """
    spec = fm.get("infrastructure") or {}
    rows: list[dict[str, Any]] = []

    for entry in spec.get("named_rows", []):
        token = entry["token"]
        if ctx.num(token):
            rows.append({"label": _label_for(ctx, token),
                         "formatted": ctx.formatted.get(token, ""),
                         "description": entry["description"]})

    other = spec.get("other") or {}
    if other and ctx.num(other["token"]):
        rows.append({"label": other["label"],
                     "formatted": ctx.formatted.get(other["token"], ""),
                     "description": other["description"]})
    return rows


def _label_for(ctx: Context, token: str) -> str:
    for row in ctx.raw.get("_cost_rows") or []:
        if row["token"] == token:
            return row["label"]
    return token.replace("cost_", "").replace("_", " ").title()


def _equipment_rows(ctx: Context) -> list[dict[str, Any]]:
    """Section 6's schedule, one row per line item on the INPUT SHEET."""
    rows = []
    for item in ctx.tables.get("equipment_line_items") or []:
        category = str(item.get("category") or "")
        is_l2 = category.startswith("_L2")
        rating = ctx.num("nameplate_l2_kw") if is_l2 else 0.0
        modeled = ctx.num("modeled_rating_l2") if is_l2 else ctx.num("modeled_rating_l3")
        rows.append({
            "description": item.get("description") or item.get("sku"),
            "sku": item.get("sku"),
            "qty": item.get("qty"),
            "nameplate": f"{rating:g} kW each" if rating else "As scheduled",
            "modeled": f"{modeled:,.2f} kW" if modeled else "N/A",
            "use": ("Longer-dwell customer charging" if is_l2
                    else "High-turnover fast charging"),
        })
    return rows


def _roi_rows(ctx: Context) -> list[dict[str, str]]:
    """Section 7's ROI table. The ITC row is present only when non-zero."""
    rows = [
        ("Total Costs Upfront", "roi_total_costs_upfront"),
        ("Carbon Credits", "roi_carbon_credits"),
    ]
    if ctx.raw.get("has_itc"):
        rows.append(("Federal ITC (30%)", "roi_itc"))
    rows += [("EVSE Revenues", "roi_evse_revenues"),
             ("Net Revenues", "roi_net_revenues")]
    return [{"label": label, "formatted": ctx.formatted.get(token, "")}
            for label, token in rows]


# --------------------------------------------------------------------------
# post-render verification
# --------------------------------------------------------------------------


def verify_output(path: Path) -> list[str]:
    """Nothing unrendered, no Excel errors, no previous client, anywhere.

    Checked on the saved file rather than the in-memory document, because the
    headers, footers and docProps are separate parts and are exactly where a
    leftover survives.
    """
    from tools.docx_edit import audit_identity, find_leftover_tokens

    problems: list[str] = []
    for part, tokens in find_leftover_tokens(path).items():
        problems.append(f"{part} still contains {', '.join(tokens)}")
    for part, names in audit_identity(path).items():
        problems.append(f"{part} still mentions {', '.join(names)}")

    import zipfile

    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not re.search(r"(document|header\d+|footer\d+)\.xml$", name):
                continue
            text = z.read(name).decode("utf-8", errors="replace")
            flat = re.sub(r"<[^>]+>", "", text)
            if re.search(r"\bNone\b", flat):
                problems.append(f"{name} contains a literal 'None'")
    return problems


# --------------------------------------------------------------------------
# the pipeline
# --------------------------------------------------------------------------


def generate(
    workbook: str | Path,
    *,
    output: str | Path | None = None,
    operator_inputs_path: str | Path | None = None,
    operator_inputs: dict | None = None,
    template: str | Path | None = None,
    force: bool = False,
    progress=None,
) -> int:
    """Full pipeline. Returns a shell exit code; raises `ProposalError` on abort."""
    def say(message: str) -> None:
        if progress:
            progress(message)
        else:
            print(message)

    template_path = Path(template or PROPOSAL_TEMPLATE)
    if not template_path.exists():
        raise ProposalError(
            "The proposal template is missing.",
            detail=f"expected {template_path}. Run: python -m tools.build_template",
        )

    inputs = dict(operator_inputs or {})
    if operator_inputs_path:
        with open(operator_inputs_path, encoding="utf-8") as fh:
            inputs.update(json.load(fh))

    fm = load_field_map()
    say("Reading the workbook...")
    with engine_mod.load(workbook, fm) as lw:
        for line in lw.detection.summary_lines():
            say(f"  {line}")

        say("Extracting figures...")
        ctx = build_context(lw, fm, operator_inputs=inputs)

        plan = build_plan(
            {**ctx.raw, **{k: bool(v) for k, v in ctx.raw.items()
                           if k.startswith("has_")}},
            substitutions={
                "projection_years": ctx.raw.get("projection_years", 5),
                "loan_term_months": ctx.raw.get("loan_term_months", 60),
            },
            disabled=inputs.get("disabled_sections") or (),
        )
        for section in plan.forced_on:
            ctx.findings.warn(
                "section-cannot-be-removed",
                f"'{section.title}' was switched off but is required in every "
                "proposal, so it has been kept.",
            )
        say(f"  {len(plan.sections)} sections, {len(plan.dropped)} suppressed")

        say("Validating...")
        from .validate import qa_report, validate

        findings = validate(ctx, fm, plan)

        out_path = Path(output) if output else default_output_path(ctx)
        try:
            out_path.parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise OutputNotWritable.for_path(out_path.parent, exc) from exc
        qa_path = out_path.with_name(out_path.stem + "_QA.txt")

        if findings.has_errors() and not force:
            # The QA report is still written. An operator who has just been told
            # "no document" needs to know why without re-running anything.
            qa_path.write_text(
                qa_report(ctx, findings, plan, output_path="(not written)"),
                encoding="utf-8",
            )
            summary = "\n".join(f"  - {f.message}" for f in findings.errors)
            raise ValidationFailed(
                f"This workbook has {len(findings.errors)} problem(s) that would "
                f"put wrong numbers in the proposal, so nothing was written:\n"
                f"{summary}\n\nDetails in {qa_path}",
                findings.errors,
            )

        say("Rendering charts...")
        from .charts import render_all

        charts = render_all(lw, ctx, out_path.parent / ".charts")

        say("Rendering the document...")
        tpl = DocxTemplate(str(template_path))
        payload = build_render_context(ctx, plan, tpl, charts, fm)

        # docxtpl lists every variable in the file, including ones inside
        # `{%p if %}` blocks that will not render. A suppressed Section 3 leaves
        # `hist_total_revenue` "missing" even though nothing will ask for it, so
        # this cannot be a hard failure - it would block every proposal for a
        # site with no history. Fill the gaps with blanks and say what was blank.
        declared = set(tpl.get_undeclared_template_variables())
        missing = sorted(declared - set(payload))
        for name in missing:
            payload[name] = ""
        if missing:
            findings.warn(
                "template-values-blank",
                "These placeholders had no value and were left blank: "
                + ", ".join(missing)
                + ". That is expected for a suppressed section, and for details "
                "nobody filled in. Anything else means the template and "
                "config/field_map.yaml have drifted apart.",
            )

        # autoescape is off by default in docxtpl, which drops a bare "&" into
        # the XML and silently loses it - "EVOLV & Commissioning" renders as
        # "EVOLV  Commissioning". Every value here is plain text, never markup.
        tpl.render(payload, autoescape=True)
        try:
            tpl.save(str(out_path))
        except OSError as exc:
            raise OutputNotWritable.for_path(out_path, exc) from exc

    say("Checking the finished document...")
    problems = verify_output(out_path)
    for problem in problems:
        findings.error("post-render", problem, where=str(out_path))

    try:
        qa_path.write_text(
            qa_report(ctx, findings, plan, output_path=str(out_path), charts=charts),
            encoding="utf-8",
        )
    except OSError as exc:
        raise OutputNotWritable.for_path(qa_path, exc) from exc

    if problems and not force:
        out_path.unlink(missing_ok=True)
        raise ValidationFailed(
            "The rendered document failed its final check and was deleted:\n"
            + "\n".join(f"  - {p}" for p in problems)
            + f"\n\nDetails in {qa_path}",
            findings.errors,
        )

    say("")
    say(f"Proposal   {out_path}")
    say(f"QA report  {qa_path}")
    if findings.warnings:
        say("")
        say(f"{len(findings.warnings)} warning(s) - read the QA report before sending:")
        for finding in findings.warnings:
            say(f"  - {finding.message}")
    return 0
